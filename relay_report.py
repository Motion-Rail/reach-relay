"""v36 (relay): cable report as an Excel workbook, from what Cable View shows.

POST /api/report/cable {stem, ribbons?: [1..36]}  -> .xlsx (desktop users only)
Sheets: Summary (counts and typical values) and Fibres (one row per fibre: continuity, iOLM loss from each
end per wavelength with the both ways average, lengths, OTDR loss, last tested). Flipped uses the same rule
as Cable View: a swap to the mirror position (F1 to F12) counts as flipped only when two mirror pairs in the
ribbon agree; one swapped pair stays Crossed.
"""
from __future__ import annotations

import asyncio
import io
import re
import time
from datetime import datetime

from fastapi import APIRouter, Header, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

WORD = {"straight": "Straight", "confirmed": "Straight", "cross": "Crossed", "flip": "Flipped", "dis": "Dis.", "unres": "Not found"}


class ReportIn(BaseModel):
    stem: str
    ribbons: list[int] = []


def _uk(t) -> str:
    if not t:
        return ""
    try:
        from zoneinfo import ZoneInfo
        return datetime.fromtimestamp(float(t), ZoneInfo("Europe/London")).strftime("%Y-%m-%d %H:%M")
    except Exception:                                   # noqa: BLE001
        return datetime.utcfromtimestamp(float(t)).strftime("%Y-%m-%d %H:%M")


def _short(rtu: str) -> str:
    m = re.match(r"^RTU\d*-([A-Z0-9]+)-\d+$", str(rtu or "").upper())
    return m.group(1) if m else str(rtu or "")


def _pick(r: dict) -> dict | None:
    e, u = r.get("e2e"), r.get("uni")
    if e and u:
        return e if (e.get("t") or 0) >= (u.get("t") or 0) else u
    return e or u


def continuity(rows: list[dict]) -> dict[int, dict]:
    """fibre -> {state, found, src, t, by}, with Flipped worked out per ribbon."""
    out = {}
    for r in rows:
        x = _pick(r)
        if x:
            st = "straight" if x.get("state") == "confirmed" else x.get("state")
            out[r["f"]] = {"state": st, "found": x.get("found"), "src": "E2E" if x is r.get("e2e") else "Uni-dir",
                           "t": x.get("t"), "by": x.get("by", "")}

    def mirror(f, found):
        return bool(found) and (int(found) - 1) // 12 == (f - 1) // 12 and (int(found) - 1) % 12 + 1 == 13 - ((f - 1) % 12 + 1)
    for rb in range(1, 37):
        pairs = {min((f - 1) % 12 + 1, 13 - ((f - 1) % 12 + 1)) for f in range((rb - 1) * 12 + 1, rb * 12 + 1)
                 if f in out and out[f]["state"] == "cross" and mirror(f, out[f]["found"])}
        if len(pairs) >= 2:
            for f in range((rb - 1) * 12 + 1, rb * 12 + 1):
                if f in out and out[f]["state"] == "cross" and mirror(f, out[f]["found"]):
                    out[f]["state"] = "flip"
    return out


IOLM_CACHE: dict[str, dict] = {}     # route id -> {"t", "item"}; latest iOLM per route, kept 10 min


def latest_iolm(fms, route_id: str) -> dict | None:
    """The newest iOLM on a route with its loss per wavelength (bulk Task outputs hold only one figure)."""
    import fms_continuity as fc
    hit = IOLM_CACHE.get(str(route_id))
    if hit and time.time() - hit["t"] < 600:
        return hit["item"]
    params = {"$filter": f"metadata/AssetId eq {int(route_id)} and metadata/TestCategory eq 'Adhoc'",
              "$orderby": "metadata/TestTime desc", "$top": "8", "$skip": "0", "$select": "resultid,metadata,brief/LinkResults,brief/GlobalVerdict"}
    r = fms.get(fc.RESULTS_URL, params=params)
    r.raise_for_status()
    j = r.json()
    item = None
    for x in j.get("results", j if isinstance(j, list) else []):
        md = x.get("metadata") or {}
        if str(md.get("TestType") or "").upper() != "IOLM" or md.get("HasError"):
            continue
        lr = (x.get("brief") or {}).get("LinkResults") or {}
        loss = {}
        for q in lr.get("Results") or []:
            try:
                loss[str(q.get("Wavelength"))] = round(float(q.get("Loss")), 3)
            except Exception:                           # noqa: BLE001
                pass
        from relay_fibres import _when
        item = {"loss": loss, "len": float(lr.get("Length")) if str(lr.get("Length") or "").replace(".", "", 1).isdigit() else None,
                "verdict": (x.get("brief") or {}).get("GlobalVerdict") or "", "t": _when(md.get("TestTime"))}
        break
    IOLM_CACHE[str(route_id)] = {"t": time.time(), "item": item}
    if len(IOLM_CACHE) > 3000:
        IOLM_CACHE.pop(next(iter(IOLM_CACHE)))
    return item


def enrich(fms, ends: list[dict], rows: list[dict]) -> list[dict]:
    """Rows with iOLM loss per wavelength read from FMS for both ends (8 at a time)."""
    from concurrent.futures import ThreadPoolExecutor
    jobs = [(r["f"], k, e["routes"].get(str(r["f"]))) for r in rows for k, e in zip("ab", ends[:2]) if e["routes"].get(str(r["f"]))]
    with ThreadPoolExecutor(8) as ex:
        got = list(ex.map(lambda j: (j[0], j[1], latest_iolm(fms, j[2])), jobs))
    out = {r["f"]: {**r, "iolm": dict(r.get("iolm") or {})} for r in rows}
    for f, k, item in got:
        if item:
            cur = out[f]["iolm"].get(k) or {}
            if not cur.get("t") or (item.get("t") or 0) >= cur.get("t", 0):
                out[f]["iolm"][k] = {**cur, **item}
    return [out[r["f"]] for r in rows]


def workbook(stem: str, ends: list[dict], rows: list[dict], ribbons: list[int]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    a = _short(ends[0]["rtu"]) if ends else "A end"
    b = _short(ends[1]["rtu"]) if len(ends) > 1 else "B end"
    if ribbons:
        rows = [r for r in rows if (r["f"] - 1) // 12 + 1 in set(ribbons)]
    cont = continuity(rows)
    wls = sorted({w for r in rows for k in ("a", "b") for w in (((r.get("iolm") or {}).get(k) or {}).get("loss") or {})},
                 key=lambda w: int(w) if str(w).isdigit() else 0) or ["1550"]
    head_fill = PatternFill("solid", fgColor="1F2A44")
    head_font = Font(bold=True, color="FFFFFF")
    fills = {"straight": "D6F2E2", "flip": "D9E5FB", "cross": "FDE7C8", "dis": "FBD5D8", "unres": "E8E3F5"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Fibres"
    cols = ["Fibre", "Ribbon", "Position", "Continuity", "Lands on", "From", "Continuity tested", "By"]
    for w in wls:
        cols += [f"iOLM {w} from {a} (dB)", f"iOLM {w} from {b} (dB)", f"iOLM {w} both ways (dB)"]
    cols += [f"Length from {a} (km)", f"Length from {b} (km)", f"iOLM verdict {a}", f"iOLM verdict {b}",
             f"OTDR loss from {a} (dB)", f"OTDR loss from {b} (dB)", "Last tested"]
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        f = r["f"]
        c = cont.get(f)
        ia, ib = ((r.get("iolm") or {}).get("a") or {}), ((r.get("iolm") or {}).get("b") or {})
        oa, ob = ((r.get("otdr") or {}).get("a") or {}), ((r.get("otdr") or {}).get("b") or {})
        line = [f"F{f:03d}", (f - 1) // 12 + 1, (f - 1) % 12 + 1,
                WORD.get(c["state"], c["state"]) if c else "Not tested",
                f"F{int(c['found']):03d}" if c and c.get("found") and c["state"] in ("cross", "flip") else "",
                c["src"] if c else "", _uk(c["t"]) if c else "", (c or {}).get("by", "")]
        for w in wls:
            la, lb = (ia.get("loss") or {}).get(w), (ib.get("loss") or {}).get(w)
            line += [la, lb, round((float(la) + float(lb)) / 2, 3) if la is not None and lb is not None else None]
        lens = [x.get("len") or y.get("len") for x, y in ((ia, oa), (ib, ob))]
        times = [x.get("t") for x in (ia, ib, oa, ob, c or {}) if x and x.get("t")]
        line += [round(float(lens[0]) / 1000, 3) if lens[0] else None, round(float(lens[1]) / 1000, 3) if lens[1] else None,
                 ia.get("verdict") or "", ib.get("verdict") or "", oa.get("loss"), ob.get("loss"), _uk(max(times)) if times else ""]
        ws.append(line)
        if c and c["state"] in fills:
            ws.cell(row=ws.max_row, column=4).fill = PatternFill("solid", fgColor=fills[c["state"]])
    for i, h in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(9, min(22, len(h) * 0.9))
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    sm = wb.create_sheet("Summary", 0)
    n = len(rows)
    count = lambda k: sum(1 for r in rows if (cont.get(r["f"]) or {}).get("state") == k)  # noqa: E731
    tested = sum(1 for r in rows if r["f"] in cont)
    both = sum(1 for r in rows if all(((r.get("iolm") or {}).get(k) or (r.get("otdr") or {}).get(k)) for k in ("a", "b")))
    avg = [x for r in rows for x in [((((r.get("iolm") or {}).get("a") or {}).get("loss") or {}).get("1550"),
                                      (((r.get("iolm") or {}).get("b") or {}).get("loss") or {}).get("1550"))]
           if x[0] is not None and x[1] is not None]
    typical = sorted((float(p) + float(q)) / 2 for p, q in avg)
    lines = [("Cable", stem), ("Ends", f"{a} to {b}"),
             ("Ribbons", ", ".join(f"R{x}" for x in sorted(set(ribbons))) if ribbons else "All"), ("Fibres", n),
             ("Report made", _uk(time.time())), ("", ""),
             ("Continuity tested", f"{tested} of {n}"), ("Straight", count("straight")), ("Flipped", count("flip")),
             ("Crossed", count("cross")), ("Dis.", count("dis")), ("Not found", count("unres")), ("Not tested", n - tested), ("", ""),
             ("Loss tested from both ends", f"{both} of {n}"),
             ("Typical both ways loss 1550 (dB)", round(typical[len(typical) // 2], 3) if typical else "no reading"), ("", ""),
             ("Notes", "Flipped: a swap to the mirror position (F1 to F12) counted only when two mirror pairs in the ribbon agree."),
             ("", "Both ways = the average of the loss from each end, which cancels the one way splice errors."),
             ("", "Results are the latest found by Reach Fibre Tester: E2E and Uni-dir runs, and FMS bulk Task results.")]
    for k, v in lines:
        sm.append([k, v])
    sm["A1"].font = Font(bold=True, size=13)
    for row in sm.iter_rows(min_row=1, max_row=sm.max_row, max_col=1):
        for c in row:
            c.font = Font(bold=True)
    sm.column_dimensions["A"].width = 32
    sm.column_dimensions["B"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_router(valid_token, check_key, sessions: dict) -> APIRouter:
    router = APIRouter()
    from relay_bulk import require_desktop

    @router.post("/api/report/cable")
    async def report(body: ReportIn, x_app_key: str | None = Header(default=None), x_session: str | None = Header(default=None)):
        check_key(x_app_key)
        token = await valid_token(x_session)
        require_desktop(sessions, x_session)
        stem = body.stem.strip().upper()
        if not re.match(r"^F-[A-Z0-9]+-[A-Z0-9]+-[A-Z](-R\d+)?$", stem):
            raise HTTPException(400, "Not a cable name, for example F-RGAC-SNBC-A-R432")
        if any(not 1 <= x <= 36 for x in body.ribbons):
            raise HTTPException(400, "Ribbons are 1 to 36")
        import relay_fibres
        ends = await relay_fibres.ensure_ends(token, stem)
        b = relay_fibres.BASE.get(stem)
        if not b or time.time() - b["t"] > relay_fibres.BASE_TTL:
            data = await asyncio.to_thread(relay_fibres.build, stem, ends)
            b = relay_fibres.BASE[stem] = {"t": time.time(), "data": data}
        rows = b["data"]["rows"]
        if body.ribbons:
            rows = [r for r in rows if (r["f"] - 1) // 12 + 1 in set(body.ribbons)]
        import relay_bulk
        import fms_continuity as fc
        fms = relay_bulk.FMS_FOR(token) if relay_bulk.FMS_FOR else fc.Fms(token)
        try:
            rows = await asyncio.to_thread(enrich, fms, ends, rows)
        except Exception:                               # noqa: BLE001
            pass                                        # the report still goes out with what Cable View has
        x = await asyncio.to_thread(workbook, stem, ends, rows, body.ribbons)
        name = f"{stem} report {datetime.now().strftime('%Y-%m-%d')}" + (f" R{min(body.ribbons)}-{max(body.ribbons)}" if body.ribbons else "") + ".xlsx"
        return Response(x, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{name}"'})

    return router
