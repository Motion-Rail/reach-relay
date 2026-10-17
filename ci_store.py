"""Cable Insight store.

One folder per cable section under data/, each holding:
  pulls/        pull workbooks loaded by hand, exactly as fms_pull wrote them
  raw/A, raw/B  every FMS result pulled, one file per fibre per end (Phase 1)
  store.sqlite  the processed results, one dataset per pull

Locations are named with fms_pull's own schedule matching (Distances.xlsx), so the
names are the ones the client workbooks use.

A splice seen from both ends uses the signed average (a gainer's negative reading is
kept). A splice seen from one end only follows the "one end only" setting:
  nominal  the missing direction is the nominal value (default 0.02 dB)
  auto     the missing direction is the same direction ribbon average at that location
           when at least half the ribbon saw it there, otherwise the nominal value
  exclude  left out of Warn+ and listed as one way only
Filled values are always marked.
"""
import gzip
import hashlib
import collections
import json
import os
import re
import shutil
import sqlite3
import statistics
import sys
import time

import openpyxl

import ci_align

SPLICE_TYPES = ("Splice", "Group")
END_MARGIN_M = 100.0
MATCH_TOL_M = 30.0
DIST_OVERRIDE = None        # a Distances workbook to use instead of the one next to fms_pull (the relay)
LEARNED_FOR = None          # cable -> learned locations [{"name","distance_m","tolerance_m","type"}] (the relay)
NAMES_SAVED = None          # (section, names) -> called after a joint is named, so the relay can keep a copy
SNAPS_SAVED = None          # (section, snapshots) -> called after a snapshot is saved or removed (the relay keeps a copy)
ODF_KINDS = ("ODF", "S-ODF", "SODF", "NODF", "REB")
JOINT_TYPES = ["Joint", "Temp joint", "Spur joint", "ODF", "S-ODF", "REB", "Other"]
FIBRES_PER_RIBBON = 12
DEFAULT_SETTINGS = {"one_end": "nominal", "nominal": 0.02, "threshold": 0.15, "split": False,
                    "sodf_limit": 0.30, "nr_ribbons": "1-18", "spur_gap": 1500, "group_tol": 100,
                    "odf_limit": 0.50}

SCHEMA = """
CREATE TABLE IF NOT EXISTS datasets(
  id INTEGER PRIMARY KEY, label TEXT, source TEXT, sha1 TEXT, imported_at TEXT,
  section TEXT, cable TEXT, site_a TEXT, site_b TEXT, t_from TEXT, t_to TEXT,
  n_fibres INTEGER, n_events INTEGER, schedule TEXT, content_sig TEXT, notes TEXT, align INTEGER);
CREATE TABLE IF NOT EXISTS fibres(
  ds INTEGER, fno INTEGER, fid TEXT, ribbon INTEGER,
  date_a TEXT, date_b TEXT, len_a REAL, len_b REAL,
  a1310 REAL, a1550 REAL, a1625 REAL, b1310 REAL, b1550 REAL, b1625 REAL);
CREATE TABLE IF NOT EXISTS events(
  ds INTEGER, fno INTEGER, ribbon INTEGER, pos REAL, type TEXT,
  ab REAL, ba REAL, dirs INTEGER, loc TEXT, bpos REAL, spur_d REAL, spur_len REAL);
CREATE TABLE IF NOT EXISTS locations(
  ds INTEGER, loc TEXT, ord INTEGER, pos REAL, kind TEXT, n_fibres INTEGER,
  scheduled INTEGER, spur_d REAL, spur_len REAL, was TEXT);
CREATE TABLE IF NOT EXISTS odf(
  ds INTEGER, fno INTEGER, ribbon INTEGER, end_ TEXT, pos REAL, type TEXT,
  refl REAL, w1310 REAL, w1550 REAL, w1625 REAL, sw REAL);
CREATE INDEX IF NOT EXISTS ev_ds ON events(ds, loc);
CREATE INDEX IF NOT EXISTS fi_ds ON fibres(ds);
"""


# --------------------------------------------------------------------------- helpers
def _f(v):
    try:
        if v is None or v == "":
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def fibre_no(fid):
    m = re.search(r"F(\d+)\s*$", str(fid or ""))
    return int(m.group(1)) if m else None


def ribbon_of(fno):
    return (fno - 1) // FIBRES_PER_RIBBON + 1 if fno else None


def section_of(fid):
    """F-RGAC-SNBC-A-R432-F001 -> RGAC-SNBC-A-R432"""
    m = re.match(r"F-(.+)-F\d+\s*$", str(fid or ""))
    return m.group(1) if m else "UNKNOWN"


def cable_of(fid):
    """F-RGAC-SNBC-A-R432-F001 -> F-RGAC-SNBC (the Distances.xlsx CABLE tag)"""
    m = re.match(r"(F-[A-Z0-9]+-[A-Z0-9]+)", str(fid or ""))
    return m.group(1) if m else ""


def sha1_of(path):
    h = hashlib.sha1()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def content_sig(rawA, rawB):
    """Fingerprint of WHICH tests a pull holds: every fibre's test time at each end.
    Two pulls with the same fingerprint hold the same results, whatever the file."""
    parts = [f"A{n}:{rawA[n]['date'][:19]}" for n in sorted(rawA) if rawA[n].get("date")]
    parts += [f"B{n}:{rawB[n]['date'][:19]}" for n in sorted(rawB) if rawB[n].get("date")]
    return hashlib.sha1("|".join(parts).encode()).hexdigest()


def is_pull_workbook(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ok = "All Events" in wb.sheetnames and "Raw A-B (End A)" in wb.sheetnames
        wb.close()
        return ok
    except Exception:
        return False


def uk_time(iso):
    """FMS times are UTC; show them as UK clock time (BST from the last Sunday of March to the last Sunday
    of October, 01:00 UTC), e.g. 14 Oct 2026 15:39. No time zone database needed (Windows has none)."""
    import datetime as _dt
    try:
        t = _dt.datetime.fromisoformat(str(iso)[:19])
    except ValueError:
        return str(iso or "")
    def last_sun(y, m):
        d = _dt.datetime(y, m, 31 if m in (3, 10) else 30, 1)
        return d - _dt.timedelta(days=(d.weekday() + 1) % 7)
    if last_sun(t.year, 3) <= t < last_sun(t.year, 10):
        t += _dt.timedelta(hours=1)
    return t.strftime("%d %b %Y %H:%M").lstrip("0")


def tested_span(t_from, t_to):
    """Start = first trace used, end = last trace used, both in UK time."""
    return f"start {uk_time(t_from)} (first trace), end {uk_time(t_to)} (last trace), UK time"


def tested_label(t_from, t_to):
    def d(s):
        try:
            return time.strptime(s[:10], "%Y-%m-%d")
        except (ValueError, TypeError):
            return None
    a, b = d(t_from), d(t_to)
    if not b:
        return "Pull"
    if a and a[:3] != b[:3]:
        if a[:2] == b[:2]:
            return f"Tests {a.tm_mday} to {time.strftime('%d %b %Y', b).lstrip('0')}"
        return f"Tests {time.strftime('%d %b', a).lstrip('0')} to {time.strftime('%d %b %Y', b).lstrip('0')}"
    return f"Tests {time.strftime('%d %b %Y', b).lstrip('0')}"


def splice_value(ab, ba, st, ravg_ab=None, ravg_ba=None):
    """(value, how, gainer). how = both | nominal | ribbon | one_way (left out) | None."""
    if ab is not None and ba is not None:
        return (ab + ba) / 2.0, "both", (ab < 0 or ba < 0)
    if ab is None and ba is None:
        return None, None, False
    mode, nom = st.get("one_end", "nominal"), float(st.get("nominal", 0.02))
    if mode == "exclude":
        return (ab if ab is not None else ba), "one_way", False
    how = "nominal"
    if ab is None:
        fill = ravg_ab if (mode == "auto" and ravg_ab is not None) else nom
        if mode == "auto" and ravg_ab is not None:
            how = "ribbon"
        return (fill + ba) / 2.0, how, False
    fill = ravg_ba if (mode == "auto" and ravg_ba is not None) else nom
    if mode == "auto" and ravg_ba is not None:
        how = "ribbon"
    return (ab + fill) / 2.0, how, False


def ribbon_averages(group, size):
    """Same direction averages at one location for one ribbon, only when at least half
    the ribbon's fibres saw the event in that direction."""
    out = []
    for d in ("ab", "ba"):
        vals = [g[d] for g in group if g[d] is not None]
        n = len({g["fno"] for g in group if g[d] is not None})
        out.append(sum(vals) / len(vals) if vals and n >= size / 2.0 else None)
    return out


# --------------------------------------------------------------------------- reading
def read_pull_workbook(path):
    """Read an fms_pull bidir workbook: both raw E2E sheets and All Events."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def raw(sheet):
        out = {}
        if sheet not in wb.sheetnames:
            return out
        for r in wb[sheet].iter_rows(min_row=3, values_only=True):
            fno = fibre_no(r[0])
            if fno is None:
                continue
            out[fno] = {"fid": r[0], "date": str(r[1] or ""), "l1550": _f(r[6]),
                        "len": _f(r[3]), "l1310": _f(r[5]), "l1625": _f(r[7])}
        return out

    rawA, rawB = raw("Raw A-B (End A)"), raw("Raw B-A (End B)")

    ws = wb["All Events"]
    hdr, events = None, []
    ab_col = ba_col = None
    for r in ws.iter_rows(values_only=True):
        if hdr is None:
            if r and "Fibre" in r and "Position (m)" in r:
                hdr = {str(h).strip(): i for i, h in enumerate(r) if h is not None}
                ab_col = next((k for k in hdr if k.startswith("A-B")), None)
                ba_col = next((k for k in hdr if k.startswith("B-A")), None)
            continue
        fid = r[hdr["Fibre"]]
        fno = fibre_no(fid)
        if fno is None:
            continue
        events.append({
            "fid": fid, "fno": fno, "position": _f(r[hdr["Position (m)"]]),
            "type": r[hdr["Type"]], "ab": _f(r[hdr[ab_col]]) if ab_col else None,
            "ba": _f(r[hdr[ba_col]]) if ba_col else None,
            "dirs": int(r[hdr["Dirs"]] or 0) if "Dirs" in hdr else None})

    site_a = site_b = ""
    if "Bidir E2E Loss Report" in wb.sheetnames:
        top = next(wb["Bidir E2E Loss Report"].iter_rows(max_row=1, values_only=True))
        site_a = str(top[0] or "").split("(")[0].strip()
        site_b = next((str(c).split("(")[0].strip() for c in top[1:]
                       if c and "End B" in str(c)), "")
    wb.close()
    return finish_data(rawA, rawB, events, site_a, site_b)


def finish_data(rawA, rawB, events, site_a, site_b):
    fids = [v["fid"] for v in rawA.values() if v.get("fid")] or [e["fid"] for e in events]
    dates = sorted(d[:19] for d in [v.get("date") for v in rawA.values()] +
                   [v.get("date") for v in rawB.values()] if d)
    return {"rawA": rawA, "rawB": rawB, "events": events, "site_a": site_a,
            "site_b": site_b, "section": section_of(fids[0]) if fids else "UNKNOWN",
            "cable": cable_of(fids[0]) if fids else "",
            "t_from": dates[0] if dates else "", "t_to": dates[-1] if dates else ""}


# --------------------------------------------------------------------------- naming
def _joint_types(dist_path, sheet):
    out = {}
    try:
        wb = openpyxl.load_workbook(dist_path, read_only=True, data_only=True)
        ws = wb[sheet]
        col = None
        for r in ws.iter_rows(values_only=True):
            if col is None:
                if r and "Location" in r and "Joint type" in r:
                    col = list(r).index("Joint type")
                continue
            if r[0]:
                out[str(r[0]).strip()] = str(r[col] or "").strip()
        wb.close()
    except Exception:
        pass
    return out


def load_fms_pull(fms_dir):
    if not fms_dir or not os.path.exists(os.path.join(fms_dir, "fms_pull.py")):
        return None
    if fms_dir not in sys.path:
        sys.path.insert(0, fms_dir)
    import fms_pull
    return fms_pull


def name_locations(events, cable, site_a, fms_dir, log=print, lengths=None):
    """Give each mid-route event its location name using fms_pull's own matching, or, for a
    cable with no Distances schedule, from the joints most of the cable shares (named from
    the app's learned locations where they match)."""
    fp = load_fms_pull(fms_dir)
    schedule, label, kinds = [], "none", {}
    dist = DIST_OVERRIDE or os.path.join(fms_dir or "", "Distances.xlsx")
    if fp is not None and os.path.exists(dist):
        sheet, _ = fp.pick_sheet_for_cable(dist, cable, site=site_a, verbose=False)
        if sheet:
            schedule = fp.load_locations(dist, sheet, verbose=False)
            kinds = _joint_types(dist, sheet)
            label = f"Distances.xlsx / {sheet}"
    if schedule:
        fp.map_locations(events, schedule, tol=120.0, verbose=False)
        fp.label_unmapped(events, tol=25.0, min_fibres=3, verbose=False)
    else:
        label, kinds = found_locations(events, cable, lengths)
    log(f"  locations named from {label}")
    return label, kinds


def found_locations(events, cable, lengths):
    learned = None
    if LEARNED_FOR:
        try:
            learned = LEARNED_FOR(cable)
        except Exception:                   # noqa: BLE001
            learned = None
    for e in events:
        e["location"] = None
        e.pop("unscheduled", None)
    res = ci_align.find_backbone(events, lengths, learned)
    if learned:
        label = f"learned locations ({res['learned']} of {len(res['names'])} joints named)"
    else:
        label = f"joints found from the results ({len(res['names'])})"
    return label, res["kinds"]


# --------------------------------------------------------------------------- backbone
def _is_mid(e, cable_len, lens=None):
    """A field event: past the launch and the ODF at End A, and short of the far end of its own fibre
    (the ODF at End B, which the ODF view judges at its own limit on End B's trace). The far end is
    each fibre's own length, so fibres shorter than most (another route) do not have their far end
    ODF judged as a splice."""
    p = e.get("position")
    end = (lens or {}).get(e.get("fno")) or cable_len
    return (p is not None and p > END_MARGIN_M and (end is None or p < end - END_MARGIN_M)
            and e.get("type") in ("Splice", "Group", "Connector", "Macrobend"))


def _align_mid(mid, scheduled, found=False):
    """Backbone distance for every mid route event, then the unscheduled joints clustered on it.
    found: no Distances schedule, so spur loops are placed from the ribbons' own unnamed joints."""
    res = ci_align.align(mid, scheduled, loose_loops=found)
    ci_align.label_loose(mid, scheduled)
    return res


def _insert_locations(con, ds, mid, kinds):
    """One row per location at its backbone distance. A location most of whose readings sit
    inside a spur loop is a spur location: the junction's distance plus its place on the spur."""
    groups = {}
    for e in mid:
        if e.get("location"):
            groups.setdefault(e["location"], []).append(e)
    rows = []
    for k, xs in groups.items():
        sp = [x for x in xs if x.get("spur_d") is not None]
        if len(sp) * 2 >= len(xs):
            pos = statistics.median(x["bpos"] for x in sp)
            sd, sl = statistics.median(x["spur_d"] for x in sp), statistics.median(x["spur_len"] for x in sp)
        else:
            on = [x for x in xs if x.get("spur_d") is None] or xs
            pos, sd, sl = statistics.median(x.get("bpos", x["position"]) for x in on), None, None
        rows.append([k, pos, sd, sl, xs])
    rows.sort(key=lambda r: (round(r[1]), -1 if r[2] is None else r[2]))
    con.executemany("INSERT INTO locations(ds,loc,ord,pos,kind,n_fibres,scheduled,spur_d,spur_len) "
                    "VALUES(?,?,?,?,?,?,?,?,?)", [
        (ds, k, i, pos, kinds.get(k, ""), len({x["fno"] for x in xs}),
         0 if any(x.get("unscheduled") for x in xs) else 1, sd, sl)
        for i, (k, pos, sd, sl, xs) in enumerate(rows)])
    return len(rows)


# --------------------------------------------------------------------------- app names
NAME_POS_M = 25.0           # a named joint and a joint in a later test are the same within this
NAME_SPUR_M = 15.0          # and, on a spur, within this along the spur


def _renameable(loc_row):
    loc = str(loc_row["loc"])
    return loc_row["scheduled"] != 1 or loc.startswith("Unnamed ")


def _find_name(names, row):
    """The saved name for a location row, matched by place."""
    best = None
    for n in names:
        if (n.get("spur_d") is None) != (row.get("spur_d") is None):
            continue
        gap = abs(n["bpos"] - row["pos"])
        if gap > NAME_POS_M:
            continue
        if n.get("spur_d") is not None and abs(n["spur_d"] - row["spur_d"]) > NAME_SPUR_M:
            continue
        if best is None or gap < best[0]:
            best = (gap, n)
    return best[1] if best else None


def apply_names(con, ds, names):
    """Give the joints of one pull the names saved in the app, by place. Returns how many."""
    if not names:
        return 0
    rows = [dict(r) for r in con.execute("SELECT * FROM locations WHERE ds=?", (ds,))]
    taken = {r["loc"] for r in rows}
    done = 0
    for n in names:
        cands = [r for r in rows if _renameable(r) and _find_name([n], r) is n]
        if not cands:
            continue
        r = min(cands, key=lambda x: abs(x["pos"] - n["bpos"]))
        if r["loc"] == n["name"]:
            con.execute("UPDATE locations SET kind=? WHERE ds=? AND loc=?", (n.get("type") or "", ds, r["loc"]))
            continue
        if n["name"] in taken:
            continue
        con.execute("UPDATE events SET loc=? WHERE ds=? AND loc=?", (n["name"], ds, r["loc"]))
        con.execute("UPDATE locations SET loc=?, kind=?, scheduled=2, was=COALESCE(was, ?) WHERE ds=? AND loc=?",
                    (n["name"], n.get("type") or "", r["loc"], ds, r["loc"]))
        taken.discard(r["loc"]); taken.add(n["name"])
        r["loc"] = n["name"]
        done += 1
    return done


def _unname(con, ds, name):
    r = con.execute("SELECT was FROM locations WHERE ds=? AND loc=? AND scheduled=2", (ds, name)).fetchone()
    if not r or not r[0]:
        return
    was = r[0]
    con.execute("UPDATE events SET loc=? WHERE ds=? AND loc=?", (was, ds, name))
    con.execute("UPDATE locations SET loc=?, kind='', scheduled=?, was=NULL WHERE ds=? AND loc=?",
                (was, 1 if was.startswith("Unnamed ") else 0, ds, name))


TRACK_STATUS = ("new", "back", "still", "fixed", "not_seen")


def _track_status(seq, order):
    """What happened to one splice over the tests (True fail, False pass, None not read), and since
    which test: new (passed before, fails now), back (failed, passed, fails again), still (fails in
    every test it was read in), fixed (failed, passes now), not_seen (no reading in the latest test)."""
    if seq[-1] is None:
        last = max(i for i, x in enumerate(seq) if x is not None)
        return "not_seen", order[last]
    run = len(seq) - 1                      # start of the latest unbroken run of the same result
    while run > 0 and (seq[run - 1] == seq[-1] or seq[run - 1] is None):
        run -= 1
    while seq[run] is None:
        run += 1
    before = [x for x in seq[:run] if x is not None]
    if seq[-1]:
        if not before:
            return "still", order[run]
        return ("back" if any(before) else "new"), order[run]
    return "fixed", order[run]


def _is_odf_loc(l):
    k = re.sub(r"[^A-Z]", "", str(l.get("kind") or "").upper())
    return k in {re.sub(r"[^A-Z]", "", x) for x in ODF_KINDS} or bool(re.search(r"\b(REB|ODF)\b", str(l.get("loc") or "").upper()))


def _lj(l):
    return {"loc": l["loc"], "pos": round(l["pos"], 1), "spur": l.get("spur"), "n_fibres": len(l.get("fnos") or ())}


def _is_named(l):
    loc = str(l.get("loc") or "")
    return not ("(unscheduled)" in loc or loc.startswith("Unnamed ") or loc.startswith("Spur at "))


def _trend_joint(joints, l, used=()):
    """The joint already in the trend that this pull's location is: same name, or for unnamed
    joints the same place (and the same place on its spur)."""
    best = None
    for j in joints:
        if j["loc"] == l["loc"]:
            return j
        if id(j) in used:
            continue                        # one location of a pull per joint
        if (j["spur_d"] is None) != (l.get("spur_d") is None):
            continue
        if _is_named(l) and j["named"]:
            continue
        gap = abs(j["pos"] - l["pos"])
        if gap > NAME_POS_M:
            continue
        if j["spur_d"] is not None and abs(j["spur_d"] - l["spur_d"]) > NAME_SPUR_M:
            continue
        if best is None or gap < best[0]:
            best = (gap, j)
    return best[1] if best else None


# --------------------------------------------------------------------------- store
class Store:
    def __init__(self, data_dir):
        self.data_dir = data_dir
        os.makedirs(data_dir, exist_ok=True)

    def sections(self):
        return sorted(d for d in os.listdir(self.data_dir)
                      if os.path.exists(os.path.join(self.data_dir, d, "store.sqlite")))

    def folder(self, section):
        return os.path.join(self.data_dir, section)

    def db(self, section):
        folder = self.folder(section)
        os.makedirs(os.path.join(folder, "pulls"), exist_ok=True)
        path = os.path.join(folder, "store.sqlite")
        keep = os.path.join(folder, f"store.before-align-v{ci_align.ALIGN_V}.sqlite")
        if os.path.exists(path) and not os.path.exists(keep):
            # a store from before this alignment is kept exactly as it was, once
            old = sqlite3.connect(path)
            try:
                cols = {r[1] for r in old.execute("PRAGMA table_info(datasets)")}
                if "align" not in cols or old.execute("SELECT 1 FROM datasets WHERE align IS NULL OR align<?",
                                                      (ci_align.ALIGN_V,)).fetchone():
                    old.execute("VACUUM INTO ?", (keep,))
            finally:
                old.close()
        con = sqlite3.connect(path)
        con.row_factory = sqlite3.Row
        con.executescript(SCHEMA)
        cols = {r[1] for r in con.execute("PRAGMA table_info(datasets)")}
        for c in ("content_sig", "notes"):
            if c not in cols:
                con.execute(f"ALTER TABLE datasets ADD COLUMN {c} TEXT")
        if "align" not in cols:
            con.execute("ALTER TABLE datasets ADD COLUMN align INTEGER")
        for tbl, add in (("events", ("bpos", "spur_d", "spur_len")), ("locations", ("spur_d", "spur_len"))):
            have = {r[1] for r in con.execute(f"PRAGMA table_info({tbl})")}
            for c in add:
                if c not in have:
                    con.execute(f"ALTER TABLE {tbl} ADD COLUMN {c} REAL")
        if "was" not in {r[1] for r in con.execute("PRAGMA table_info(locations)")}:
            con.execute("ALTER TABLE locations ADD COLUMN was TEXT")
        # upgrade pulls stored by v0.1: fingerprint them and give them the new names
        for d in con.execute("SELECT id, label, t_from, t_to FROM datasets WHERE content_sig IS NULL").fetchall():
            rows = con.execute("SELECT fno, date_a, date_b FROM fibres WHERE ds=?", (d["id"],)).fetchall()
            sig = content_sig({r["fno"]: {"date": r["date_a"]} for r in rows},
                              {r["fno"]: {"date": r["date_b"]} for r in rows})
            label = d["label"]
            if re.match(r"^Pull \d{1,2} \w{3} \d{4}$", label or ""):
                label = tested_label(d["t_from"], d["t_to"])
            con.execute("UPDATE datasets SET content_sig=?, label=?, notes=COALESCE(notes,'{}') WHERE id=?",
                        (sig, label, d["id"]))
        con.commit()
        todo = con.execute("SELECT id FROM datasets WHERE align IS NULL OR align<?", (ci_align.ALIGN_V,)).fetchall()
        if todo:
            for d in todo:
                self._realign(con, d["id"])
        return con

    def _realign(self, con, ds):
        """Put a stored pull on the backbone scale: pulls stored before alignment, or by an
        older version of it. Only the derived columns change (backbone distance, spur place and
        the unscheduled joint labels); the results themselves are never touched."""
        lens = [r[0] for r in con.execute("SELECT len_a FROM fibres WHERE ds=? AND len_a", (ds,))]
        cable_len = statistics.median(lens) if lens else None
        locs = {r["loc"]: dict(r) for r in con.execute("SELECT * FROM locations WHERE ds=?", (ds,))}
        sched = {k for k, v in locs.items() if v["scheduled"] == 1}       # 2: named in the app, matched again by place
        info = con.execute("SELECT cable, schedule, section FROM datasets WHERE id=?", (ds,)).fetchone()
        own = str(info["schedule"] or "").startswith("Distances")
        evs = []
        for r in con.execute("SELECT rowid, fno, pos, type, loc FROM events WHERE ds=? ORDER BY rowid", (ds,)):
            evs.append({"rowid": r["rowid"], "fno": r["fno"], "position": r["pos"], "type": r["type"],
                        "location": r["loc"] if r["loc"] in sched else None, "loc_was": r["loc"]})
        flen = {r[0]: r[1] or r[2] for r in con.execute("SELECT fno, len_a, len_b FROM fibres WHERE ds=?", (ds,))}
        mid = [e for e in evs if _is_mid(e, cable_len, flen)]
        kinds = {k: v["kind"] for k, v in locs.items()}
        if not own:
            # no Distances schedule: the shared joints are found (and named from learned locations)
            lens_by = {r[0]: r[1] or r[2] for r in con.execute("SELECT fno, len_a, len_b FROM fibres WHERE ds=?", (ds,))}
            label, kinds = found_locations(mid, info["cable"], lens_by)
            sched = {e["location"] for e in mid if e.get("location")}
            con.execute("UPDATE datasets SET schedule=? WHERE id=?", (label, ds))
        _align_mid(mid, sched, found=not own)
        for e in evs:
            if "bpos" not in e:             # launch, ODFs and fibre ends: never a joint
                e.update(bpos=e["position"], spur_d=None, spur_len=None, location=None)
        con.executemany("UPDATE events SET loc=?, bpos=?, spur_d=?, spur_len=? WHERE rowid=?",
                        [(e["location"], e["bpos"], e["spur_d"], e["spur_len"], e["rowid"]) for e in evs])
        con.execute("DELETE FROM locations WHERE ds=?", (ds,))
        _insert_locations(con, ds, mid, kinds)
        apply_names(con, ds, self.names(info["section"]))
        con.execute("UPDATE datasets SET align=? WHERE id=?", (ci_align.ALIGN_V, ds))
        con.commit()

    # ---------------------------------------------------------------- assets and their jointers (v0.12)
    def assets(self, section):
        """Joints and ODFs as assets with who made them, from the installation tracker:
        assets.json {"source", "assets": [{key, name, asset, type, mileage, jointers, completed, qcs,
        match: [joint names], bpos: [joint places], at: [{ribbons, pos, role}], ribbons, sodf}]}.
        An asset is found by its joint name in the schedule (match), by the place of a joint found from
        the results (bpos, 30 m), or where each ribbon group reads it on its own trace (at, raw distance)."""
        try:
            with open(os.path.join(self.folder(section), "assets.json"), encoding="utf-8") as fh:
                return json.load(fh)
        except (OSError, ValueError):
            return {"assets": []}

    def save_assets(self, section, data):
        os.makedirs(self.folder(section), exist_ok=True)
        path = os.path.join(self.folder(section), "assets.json")
        with open(path + ".tmp", "w", encoding="utf-8") as fh:
            json.dump(data, fh, indent=1)
        os.replace(path + ".tmp", path)

    def remediations(self, section):
        """Remediation tracker rows for this cable (remediations.json, from the cable's Remediation Tracker):
        {"source", "rows": [{joint, part, loc, kind (splice|odf), ribbon, fno, status, final, exc,
        attempts: [{date, tech, result, pre, post, notes}]}]}. joint is the asset key in assets.json."""
        try:
            with open(os.path.join(self.folder(section), "remediations.json"), encoding="utf-8") as fh:
                return json.load(fh)
        except (OSError, ValueError):
            return {"rows": []}

    @staticmethod
    def _rem_tally(R, assets, same):
        """Remediations against the original jointer (ribbons re-spliced at their joints) and against the
        technician who did each attempt (first attempt pass rate, exceptions left)."""
        jt = {a["key"]: (a.get("jointers") or "") for a in assets}
        rows = [r for r in R.get("rows", []) if str(r.get("status") or "").upper() != "SUPERSEDED"]
        per_joint = {}
        for r in rows:
            j = per_joint.setdefault(r["joint"], {"ribbons": set(), "fibres": set(), "exceptions": 0, "open": 0})
            j["ribbons"].add(r["ribbon"])
            j["fibres"].add(r["fno"])
            st_ = str(r.get("status") or "").upper()
            j["exceptions"] += st_.startswith("EXCEPTION")
            j["open"] += st_ in ("IN PROGRESS", "NOT STARTED", "")
        tech = {}
        for r in rows:
            done = [a for a in r.get("attempts", []) if str(a.get("result") or "").lower() in ("pass", "fail")]
            for i, a in enumerate(done):
                who = re.split(r"[\s/,-]+", (a.get("tech") or "").strip().upper())[0] or "none"
                who = same.get(who, who)
                t = tech.setdefault(who, {"tech": who, "attempts": 0, "passed": 0, "first": 0, "first_pass": 0,
                                          "ribbons": set(), "exceptions": 0, "odf": 0})
                ok = str(a["result"]).lower() == "pass"
                t["attempts"] += 1
                t["passed"] += ok
                t["ribbons"].add((r["joint"], r["ribbon"]))
                t["odf"] += r.get("kind") == "odf"
                if i == 0:
                    t["first"] += 1
                    t["first_pass"] += ok
                if i == len(done) - 1 and str(r.get("status") or "").upper().startswith("EXCEPTION"):
                    t["exceptions"] += 1
        techs = []
        for t in tech.values():
            t["ribbons"] = len(t["ribbons"])
            t["first_pct"] = round(100.0 * t["first_pass"] / t["first"], 1) if t["first"] else None
            t["pass_pct"] = round(100.0 * t["passed"] / t["attempts"], 1) if t["attempts"] else None
            techs.append(t)
        techs.sort(key=lambda t: -t["attempts"])
        joints = {k: {"ribbons_remediated": len(v["ribbons"]), "fibres_remediated": len(v["fibres"]),
                      "exceptions": v["exceptions"], "open": v["open"],
                      "remediated_list": ", ".join(f"R{x}" for x in sorted(v["ribbons"])), "jointers": jt.get(k, "not in assets")}
                  for k, v in per_joint.items()}
        return joints, techs

    def jointer_tally(self, section, ds, st):
        """Out of spec ribbons at every joint, tallied against the jointer who made it. A ribbon is out of
        spec at a joint when any of its fibres fails there (its worst splice, judged as in the cable view;
        one end only readings left out when the settings leave them out). The whole ribbon is re-spliced."""
        con = self.db(section)
        A = self.assets(section)
        an = self.analysis(section, ds, st)
        nr = an["nr"]
        tested = {r["fno"]: r["ribbon"] for r in con.execute("SELECT fno, ribbon FROM fibres WHERE ds=?", (ds,))}
        sizes = collections.Counter(tested.values())
        locpos = {r["loc"]: r["pos"] for r in con.execute("SELECT loc, pos FROM locations WHERE ds=?", (ds,))}
        evs = [dict(r) for r in con.execute(
            "SELECT fno, ribbon, loc, ab, ba, pos FROM events WHERE ds=? AND type IN ('Splice','Group')", (ds,))]
        thr = float(st["threshold"])
        joints, used = [], set()
        same = {k.upper(): v.upper() for k, v in (A.get("same") or {}).items()}
        rem_j, techs = self._rem_tally(self.remediations(section), A.get("assets", []), same)
        for a in A.get("assets", []):
            if a.get("odf_end"):
                continue
            mine = []
            for i, e in enumerate(evs):
                loc = e["loc"] or ""
                hit = any(loc == m or loc.startswith(m + " ") for m in a.get("match") or [])
                if not hit and e["loc"] and a.get("bpos"):
                    lp = locpos.get(e["loc"])
                    hit = lp is not None and any(abs(lp - b) <= 30 for b in a["bpos"])
                if not hit:
                    for w in a.get("at") or []:
                        if e["ribbon"] in parse_ribbons(w["ribbons"]) and abs(e["pos"] - w["pos"]) <= w.get("tol", 30):
                            hit = True
                            break
                if hit:
                    mine.append(e)
                    used.add(i)
            per = collections.defaultdict(list)
            for e in mine:
                per[e["ribbon"]].append(e)
            spliced = parse_ribbons(a["ribbons"]) if a.get("ribbons") else {r for r, g in per.items() if len(g) >= 3}
            spliced &= set(sizes)
            out, fibres_fail, worst, n_spl = [], 0, None, 0
            for rb in sorted(spliced):
                g = per.get(rb, [])
                rav = ribbon_averages(g, sizes.get(rb, FIBRES_PER_RIBBON))
                lim = float(st.get("sodf_limit", 0.30)) if a.get("sodf") and (not nr or rb in nr) else thr
                byf = collections.defaultdict(list)
                for e in g:
                    byf[e["fno"]].append(e)
                bad = []
                for f, es in byf.items():
                    vals = [splice_value(e["ab"], e["ba"], st, *rav) for e in es]
                    vals = [(v, h) for v, h, _ in vals if v is not None and h != "one_way"]
                    n_spl += len(vals)
                    if not vals:
                        continue
                    v = max(x[0] for x in vals)
                    worst = v if worst is None or v > worst else worst
                    if v >= lim:
                        bad.append((f, round(v, 3)))
                if bad:
                    out.append({"ribbon": rb, "fibres": len(bad), "worst": max(b[1] for b in bad),
                                "fails": [f"F{f:03d} {v:.2f}" for f, v in sorted(bad)]})
                    fibres_fail += len(bad)
            joints.append({k: a.get(k, "") for k in ("key", "name", "asset", "type", "mileage", "jointers", "completed", "qcs")}
                          | {k: v for k, v in rem_j.get(a.get("key"), {}).items() if k != "jointers"}
                          | {"ribbons_spliced": len(spliced), "ribbons_out": len(out), "fibres_fail": fibres_fail,
                             "worst": round(worst, 3) if worst is not None else None, "readings": n_spl,
                             "out": out, "out_list": ", ".join(f"R{o['ribbon']}" for o in out)})

        def tally(key_of):
            t = {}
            for j in joints:
                for who in key_of(j):
                    x = t.setdefault(who, {"jointer": who, "joints": 0, "ribbons_spliced": 0, "ribbons_out": 0,
                                           "fibres_fail": 0, "joints_with_out": 0, "joint_list": []})
                    x["joints"] += 1
                    x["ribbons_spliced"] += j["ribbons_spliced"]
                    x["ribbons_out"] += j["ribbons_out"]
                    x["fibres_fail"] += j["fibres_fail"]
                    x["joints_with_out"] += bool(j["ribbons_out"])
                    x["joint_list"].append(f"{j['key']} {j['ribbons_out']}")
                    x["ribbons_remediated"] = x.get("ribbons_remediated", 0) + j.get("ribbons_remediated", 0)
                    x["exceptions"] = x.get("exceptions", 0) + j.get("exceptions", 0)
            for x in t.values():
                x["per_joint"] = round(x["ribbons_out"] / x["joints"], 2) if x["joints"] else None
                x["pct"] = round(100.0 * x["ribbons_out"] / x["ribbons_spliced"], 1) if x["ribbons_spliced"] else None
            return sorted(t.values(), key=lambda x: (-x["ribbons_out"], x["jointer"]))
        recorded = lambda j: [] if not j["jointers"] or j["jointers"].startswith("not") else [j["jointers"].replace(" ", "")]   # noqa: E731
        people = lambda j: list(dict.fromkeys(same.get(p.strip().upper(), p.strip().upper()) for p in recorded(j)[0].split("/"))) if recorded(j) else []   # noqa: E731
        nobody = [j for j in joints if not recorded(j)]
        info = dict(con.execute("SELECT id, label, t_from, t_to FROM datasets WHERE id=?", (ds,)).fetchone())
        # ODF ends (no splice tally): remediations only
        odfs = []
        for a in A.get("assets", []):
            if a.get("odf_end"):
                odfs.append({k: a.get(k, "") for k in ("key", "name", "type", "jointers", "completed")} | rem_j.get(a["key"], {}))
        R = self.remediations(section)
        return {"dataset": info, "settings": st, "source": A.get("source", ""), "joints": joints,
                "people": A.get("people", {}), "technicians": techs, "odf_ends": odfs, "rem_source": R.get("source", ""),
                "rem_unmatched": sorted({r["loc"] for r in R.get("rows", []) if r["joint"] not in {a["key"] for a in A.get("assets", [])}}),
                "by_jointer": tally(lambda j: recorded(j)), "by_person": tally(people),
                "not_recorded": {"joints": len(nobody), "ribbons_out": sum(j["ribbons_out"] for j in nobody)},
                "totals": {"joints": len(joints), "ribbons_spliced": sum(j["ribbons_spliced"] for j in joints),
                           "ribbons_out": sum(j["ribbons_out"] for j in joints),
                           "fibres_fail": sum(j["fibres_fail"] for j in joints)}}

    # ---------------------------------------------------------------- names given in the app
    def names(self, section):
        """Joints named in the app for this cable: [{"name","type","bpos","spur_d","spur_len","by","at"}].
        Each is matched by place (backbone distance, and place on its spur), so every later test,
        and every pull already stored, takes the name."""
        if not section:
            return []
        try:
            with open(os.path.join(self.folder(section), "names.json"), encoding="utf-8") as fh:
                got = json.load(fh)
            return [n for n in got.get("names", []) if n.get("name") and n.get("bpos") is not None]
        except (OSError, ValueError):
            return []

    def save_names(self, section, names):
        os.makedirs(self.folder(section), exist_ok=True)
        path = os.path.join(self.folder(section), "names.json")
        tmp = path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump({"names": names}, fh, indent=1)
        os.replace(tmp, path)
        if NAMES_SAVED:
            try:
                NAMES_SAVED(section, names)
            except Exception:               # noqa: BLE001
                pass

    def set_name(self, section, ds, loc, name, typ="", by=""):
        """Name a joint (or rename one named in the app); an empty name takes the name off.
        Joints named by the Distances schedule are not changed here."""
        con = self.db(section)
        row = con.execute("SELECT * FROM locations WHERE ds=? AND loc=?", (ds, loc)).fetchone()
        if not row:
            raise KeyError("No such joint in this pull")
        row = dict(row)
        if row["scheduled"] == 1 and not str(loc).startswith("Unnamed "):
            raise ValueError("This joint is named by the Distances schedule")
        name = " ".join(str(name or "").split())[:60]
        typ = typ if typ in JOINT_TYPES else ""
        names = self.names(section)
        here = _find_name(names, row)
        if name:
            clash = next((n for n in names if n["name"].lower() == name.lower() and n is not here), None)
            if clash or any(r[0].lower() == name.lower() and r[0] != loc for r in
                            con.execute("SELECT loc FROM locations WHERE ds=? AND scheduled=1", (ds,))):
                raise ValueError(f"{name} is already used on this cable")
            entry = {"name": name, "type": typ, "bpos": round(row["pos"], 1),
                     "spur_d": None if row["spur_d"] is None else round(row["spur_d"], 1),
                     "spur_len": None if row["spur_len"] is None else round(row["spur_len"], 1),
                     "by": by, "at": time.strftime("%Y-%m-%d %H:%M")}
            if here:
                names[names.index(here)] = entry
            else:
                names.append(entry)
        elif here:
            names.remove(here)
        names.sort(key=lambda n: (n["bpos"], -1 if n["spur_d"] is None else n["spur_d"]))
        self.save_names(section, names)
        for d in [r[0] for r in con.execute("SELECT id FROM datasets")]:
            if not name and here:
                _unname(con, d, here["name"])
            apply_names(con, d, names)
        con.commit()
        return {"names": names, "name": name or None}

    def find_same(self, section, sig):
        if not os.path.exists(os.path.join(self.folder(section), "store.sqlite")):
            return None
        con = self.db(section)
        r = con.execute("SELECT id, label FROM datasets WHERE content_sig=?", (sig,)).fetchone()
        return dict(r) if r else None

    # ---------------------------------------------------------------- import
    def import_workbook(self, path, fms_dir, label="", log=print):
        log(f"Reading {os.path.basename(path)} ...")
        data = read_pull_workbook(path)
        section = data["section"]
        digest = sha1_of(path)
        sig = content_sig(data["rawA"], data["rawB"])
        same = self.find_same(section, sig)
        if same:
            log(f"  same results as {same['label']}, already in the store")
            return {"section": section, "ds": same["id"], "new": False, "same_as": same["label"]}
        keep = os.path.join(self.folder(section), "pulls", os.path.basename(path))
        os.makedirs(os.path.dirname(keep), exist_ok=True)
        if os.path.abspath(keep) != os.path.abspath(path):
            if os.path.exists(keep) and sha1_of(keep) != digest:
                stem, ext = os.path.splitext(keep)
                keep = f"{stem}_{digest[:8]}{ext}"
            if not os.path.exists(keep):
                shutil.copy2(path, keep)
        ds = self.store_data(data, os.path.basename(keep), digest, sig, fms_dir, label, log)
        return {"section": section, "ds": ds, "new": True}

    def store_data(self, data, source, digest, sig, fms_dir, label="", log=print, notes=None):
        section = data["section"]
        con = self.db(section)
        lens = [v["len"] for v in data["rawA"].values() if v.get("len")]
        cable_len = statistics.median(lens) if lens else None
        flen = {n: (data["rawA"].get(n) or {}).get("len") or (data["rawB"].get(n) or {}).get("len")
                for n in set(data["rawA"]) | set(data["rawB"])}
        mid = [e for e in data["events"] if _is_mid(e, cable_len, flen)]
        lens_by = {n: (data["rawA"].get(n) or {}).get("len") or (data["rawB"].get(n) or {}).get("len")
                   for n in set(data["rawA"]) | set(data["rawB"])}
        schedule, kinds = name_locations(mid, data["cable"], data["site_a"], fms_dir, log=log, lengths=lens_by)
        res = _align_mid(mid, {e["location"] for e in mid if e.get("location") and not e.get("unscheduled")},
                         found=not schedule.startswith("Distances"))
        log(f"  backbone: {res['anchors']} shared joints, spur loops on {res['loops']} ribbon stretches")
        label = label or tested_label(data["t_from"], data["t_to"])
        cur = con.execute(
            "INSERT INTO datasets(label,source,sha1,imported_at,section,cable,site_a,site_b,"
            "t_from,t_to,n_fibres,n_events,schedule,content_sig,notes,align) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (label, source, digest, time.strftime("%Y-%m-%d %H:%M"), section,
             data["cable"], data["site_a"], data["site_b"], data["t_from"], data["t_to"],
             len(set(data["rawA"]) | set(data["rawB"])), len(data["events"]), schedule, sig,
             json.dumps(notes or {}), ci_align.ALIGN_V))
        ds = cur.lastrowid
        rows = []
        for n in sorted(set(data["rawA"]) | set(data["rawB"])):
            a, b = data["rawA"].get(n, {}), data["rawB"].get(n, {})
            rows.append((ds, n, a.get("fid") or b.get("fid"), ribbon_of(n), a.get("date"),
                         b.get("date"), a.get("len"), b.get("len"), a.get("l1310"),
                         a.get("l1550"), a.get("l1625"), b.get("l1310"), b.get("l1550"),
                         b.get("l1625")))
        con.executemany("INSERT INTO fibres VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
        con.executemany("INSERT INTO odf VALUES(?,?,?,?,?,?,?,?,?,?,?)", [
            (ds, o["fno"], ribbon_of(o["fno"]), o["end"], o.get("pos"), o.get("type"), o.get("refl"),
             o.get("w1310"), o.get("w1550"), o.get("w1625"), o.get("sw")) for o in data.get("odf") or []])
        con.executemany("INSERT INTO events(ds,fno,ribbon,pos,type,ab,ba,dirs,loc,bpos,spur_d,spur_len) "
                        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", [
            (ds, e["fno"], ribbon_of(e["fno"]), e["position"], e["type"], e["ab"], e["ba"],
             e["dirs"], e.get("location"), e.get("bpos", e["position"]), e.get("spur_d"), e.get("spur_len"))
            for e in data["events"]])
        n_locs = _insert_locations(con, ds, mid, kinds)
        named = apply_names(con, ds, self.names(section))
        if named:
            log(f"  {named} joints named in the app")
        con.commit()
        log(f"  stored: {len(rows)} fibres, {len(data['events'])} events, {n_locs} locations")
        return ds

    # ---------------------------------------------------------------- analysis
    def analysis(self, section, ds, st):
        """Fibre groups, spurs and the per location limit for one pull."""
        con = self.db(section)
        fibres = [dict(r) for r in con.execute("SELECT * FROM fibres WHERE ds=? ORDER BY fno", (ds,))]
        locs = [dict(r) for r in con.execute("SELECT * FROM locations WHERE ds=? ORDER BY ord", (ds,))]
        loc_f = {}
        for r in con.execute("SELECT fno, loc FROM events WHERE ds=? AND loc IS NOT NULL", (ds,)):
            loc_f.setdefault(r["loc"], set()).add(r["fno"])
        groups = refine_groups(fibre_groups(fibres, float(st.get("group_tol", 100))), loc_f)
        for l in locs:
            l["spur_leg"], l["spur_m"] = ci_align.spur_place(l.get("spur_d"), l.get("spur_len"))
        spurs = aligned_spurs(locs, loc_f, groups) or \
            detect_spurs(locs, loc_f, groups, len(fibres), float(st.get("spur_gap", 1500)))
        spur_of = {l: sp["name"] for sp in spurs for l in sp["locations"]}
        nr = parse_ribbons(st.get("nr_ribbons", ""))
        for l in locs:
            l["spur"] = spur_of.get(l["loc"])
            l["sodf"] = is_sodf(l.get("kind"))
            l["fnos"] = loc_f.get(l["loc"], set())
        return {"fibres": fibres, "locs": locs, "groups": groups, "spurs": spurs, "nr": nr}

    def _limit(self, loc_meta, ribbon, st, nr):
        if loc_meta and loc_meta.get("sodf") and (not nr or ribbon in nr):
            return float(st.get("sodf_limit", 0.30))
        return float(st["threshold"])

    def _groups(self, con, ds, extra="", args=()):
        evs = [dict(r) for r in con.execute(
            "SELECT fno, ribbon, loc, ab, ba, type, pos FROM events WHERE ds=? AND loc IS NOT NULL "
            "AND type IN ('Splice','Group')" + extra, (ds,) + tuple(args))]
        sizes = {r["ribbon"]: r["n"] for r in con.execute(
            "SELECT ribbon, COUNT(*) n FROM fibres WHERE ds=? GROUP BY ribbon", (ds,))}
        groups = {}
        for e in evs:
            groups.setdefault((e["ribbon"], e["loc"]), []).append(e)
        ravg = {k: ribbon_averages(g, sizes.get(k[0], FIBRES_PER_RIBBON)) for k, g in groups.items()}
        return groups, ravg

    # ---------------------------------------------------------------- reads
    def datasets(self, section):
        con = self.db(section)
        out = []
        for r in con.execute("SELECT * FROM datasets ORDER BY t_to DESC, id DESC"):
            d = dict(r)
            both = con.execute("SELECT SUM(date_a IS NOT NULL AND date_a<>'' AND date_b IS NOT NULL AND date_b<>''),"
                               " SUM((date_a IS NULL OR date_a='') <> (date_b IS NULL OR date_b=''))"
                               " FROM fibres WHERE ds=?", (d["id"],)).fetchone()
            d["both_ends"], d["one_end"] = both[0] or 0, both[1] or 0
            try:
                d["notes"] = json.loads(d.get("notes") or "{}")
            except ValueError:
                d["notes"] = {}
            out.append(d)
        return out

    def summary(self, section, ds, st, group=None):
        con = self.db(section)
        info = dict(con.execute("SELECT * FROM datasets WHERE id=?", (ds,)).fetchone())
        an = self.analysis(section, ds, st)
        gsel = next((g for g in an["groups"] if g["id"] == group), None)
        keep = gsel["fnos"] if gsel else None
        fibres = [f for f in an["fibres"] if keep is None or f["fno"] in keep]
        meta = {l["loc"]: l for l in an["locs"]}
        groups, ravg = self._groups(con, ds)
        cells, vals = {}, []
        counts = {"splices": 0, "both": 0, "filled": 0, "filled_ribbon": 0, "fail": 0,
                  "fail_both": 0, "fail_filled": 0, "fail_sodf": 0, "pass": 0, "pass_low": 0,
                  "gainers": 0, "one_way_left_out": 0, "one_way_high": 0, "spur_fail": 0}
        for key, grp in groups.items():
            m = meta.get(key[1])
            for e in grp:
                if keep is not None and e["fno"] not in keep:
                    continue
                v, how, gain = splice_value(e["ab"], e["ba"], st, *ravg[key])
                if v is None:
                    continue
                lim = self._limit(m, key[0], st, an["nr"])
                c = cells.setdefault(key, {"ribbon": key[0], "loc": key[1], "worst": None, "worst_how": None,
                                           "n": 0, "fail": 0, "filled": 0, "gainer": 0, "one_way": 0,
                                           "one_way_high": 0, "limit": lim})
                counts["splices"] += 1
                c["n"] += 1
                if how == "one_way":
                    counts["one_way_left_out"] += 1
                    c["one_way"] += 1
                    if v >= lim:
                        counts["one_way_high"] += 1
                        c["one_way_high"] += 1
                    continue
                if how == "both":
                    counts["both"] += 1
                    counts["gainers"] += gain
                    c["gainer"] += gain
                else:
                    counts["filled"] += 1
                    counts["filled_ribbon"] += how == "ribbon"
                    c["filled"] += 1
                if v >= lim:
                    counts["fail"] += 1
                    c["fail"] += 1
                    counts["fail_both" if how == "both" else "fail_filled"] += 1
                    if m and m.get("spur"):
                        counts["spur_fail"] += 1
                else:
                    counts["pass"] += 1
                    counts["pass_low"] += v < 0.10
                    if m and m.get("sodf") and lim > float(st["threshold"]) and v >= float(st["threshold"]):
                        counts["fail_sodf"] += 1      # would fail at the single splice limit
                vals.append((v, v >= lim))
                if c["worst"] is None or v > c["worst"]:
                    c["worst"], c["worst_how"] = round(v, 3), how
        nr_suggest = suggest_nr(groups, meta)
        thr = float(st["threshold"])
        edges = [("Gainer (under 0)", None, 0.0), ("0 to 0.05", 0.0, 0.05), ("0.05 to 0.10", 0.05, 0.10)]
        if thr > 0.10:
            edges.append((f"0.10 to {thr:g}", 0.10, thr))
        edges += [(f"{thr:g} to 0.30", thr, 0.30), ("0.30 to 0.50", 0.30, 0.50), ("0.50 and over", 0.50, None)]
        judged = len(vals)
        bands = []
        for lbl, lo, hi in edges:
            xs = [f for v, f in vals if (lo is None or v >= lo) and (hi is None or v < hi)]
            bands.append({"label": lbl, "lo": lo, "hi": hi, "n": len(xs), "fail": sum(xs),
                          "pct": round(100.0 * len(xs) / judged, 1) if judged else 0.0})
        pc = lambda n: round(100.0 * n / judged, 1) if judged else 0.0          # noqa: E731
        pct = {"pass": pc(counts["pass"]), "fail": pc(counts["fail"]), "pass_low": pc(counts["pass_low"]),
               "both": pc(counts["both"]), "filled": pc(counts["filled"]), "fail_sodf": pc(counts["fail_sodf"]),
               "spur_fail": pc(counts["spur_fail"]), "gainers": round(100.0 * counts["gainers"] / counts["both"], 1) if counts["both"] else 0.0,   # of the both end splices
               "judged": judged}
        e2e = [((f["a1550"] or 0) + (f["b1550"] or 0)) / 2 for f in fibres
               if f["a1550"] is not None and f["b1550"] is not None]
        nmin = max(3, round(len(fibres) * 0.05))
        locs = []
        for l in an["locs"]:
            n = len(l["fnos"] & keep) if keep is not None else len(l["fnos"])
            if n == 0:
                continue
            locs.append({k: v for k, v in l.items() if k != "fnos"} | {"n_group": n, "minor": n < nmin})
        return {"dataset": info, "locations": locs,
                "ribbons": sorted({f["ribbon"] for f in fibres}),
                "fibres_tested": len(fibres), "both_ends": sum(1 for f in fibres if f["date_a"] and f["date_b"]),
                "counts": counts, "settings": st, "bands": bands, "pct": pct,
                "e2e_avg": round(sum(e2e) / len(e2e), 3) if e2e else None,
                "cells": list(cells.values()),
                "groups": [{k: v for k, v in g.items() if k != "fnos"} for g in an["groups"]],
                "group": group if gsel else None,
                "nr_suggest": nr_suggest,
                "spurs": [sp | {"fail": sum(c["fail"] for c in cells.values() if c["loc"] in sp["locations"])}
                          for sp in an["spurs"]]}

    def location(self, section, ds, loc, ribbon, st):
        con = self.db(section)
        an = self.analysis(section, ds, st)
        meta = next((l for l in an["locs"] if l["loc"] == loc), None)
        q = ("SELECT e.fno, e.ribbon, e.pos, e.type, e.ab, e.ba, f.fid FROM events e "
             "JOIN fibres f ON f.ds=e.ds AND f.fno=e.fno WHERE e.ds=? AND e.loc=?")
        args = [ds, loc]
        if ribbon:
            q += " AND e.ribbon=?"
            args.append(int(ribbon))
        raw = [dict(r) for r in con.execute(q + " ORDER BY e.fno, e.pos", args)]
        _, ravg = self._groups(con, ds, " AND loc=?", (loc,))
        rows = []
        for r in raw:
            sp = r["type"] in SPLICE_TYPES
            v, how, gain = splice_value(r["ab"], r["ba"], st, *ravg.get((r["ribbon"], loc), (None, None)))
            lim = self._limit(meta, r["ribbon"], st, an["nr"])
            rows.append({"fno": r["fno"], "fid": r["fid"], "ribbon": r["ribbon"], "pos": r["pos"],
                         "type": r["type"], "ab": r["ab"], "ba": r["ba"], "limit": lim,
                         "value": None if v is None else round(v, 3), "how": how, "gainer": gain,
                         "splice": sp, "fail": bool(sp and how not in ("one_way", None) and v >= lim)})
        tested = sum(1 for f in an["fibres"] if not ribbon or f["ribbon"] == int(ribbon))
        loc_out = {k: v for k, v in meta.items() if k != "fnos"} if meta else {"loc": loc}
        return {"location": loc_out, "rows": rows, "tested": tested, "settings": st}

    def ribbon(self, section, ds, ribbon, st):
        con = self.db(section)
        ribbon = int(ribbon)
        an = self.analysis(section, ds, st)
        meta = {l["loc"]: l for l in an["locs"]}
        fibres = [f for f in an["fibres"] if f["ribbon"] == ribbon]
        fn = {f["fno"] for f in fibres}
        locs = [{k: v for k, v in l.items() if k != "fnos"} for l in an["locs"] if l["fnos"] & fn]
        groups, ravg = self._groups(con, ds, " AND ribbon=?", (ribbon,))
        cells = {}
        for key, grp in groups.items():
            lim = self._limit(meta.get(key[1]), ribbon, st, an["nr"])
            for e in grp:
                v, how, gain = splice_value(e["ab"], e["ba"], st, *ravg[key])
                if v is None:
                    continue
                k = f"{e['fno']}|{e['loc']}"
                rank = -1 if how == "one_way" else v
                if k not in cells or rank > cells[k]["_rank"]:
                    cells[k] = {"value": round(v, 3), "ab": e["ab"], "ba": e["ba"], "how": how,
                                "gainer": gain, "limit": lim, "_rank": rank}
        for c in cells.values():
            c.pop("_rank", None)
        return {"ribbon": ribbon, "fibres": fibres, "locations": locs, "cells": cells, "settings": st}


    # ---------------------------------------------------------------- ODF connectors
    def backfill_odf(self, section, ds, odf):
        """Pulls stored before the ODF view have no ODF rows: add them when the same results come round."""
        con = self.db(section)
        if con.execute("SELECT COUNT(*) FROM odf WHERE ds=?", (ds,)).fetchone()[0] or not odf:
            return False
        con.executemany("INSERT INTO odf VALUES(?,?,?,?,?,?,?,?,?,?,?)", [
            (ds, o["fno"], ribbon_of(o["fno"]), o["end"], o.get("pos"), o.get("type"), o.get("refl"),
             o.get("w1310"), o.get("w1550"), o.get("w1625"), o.get("sw")) for o in odf])
        con.commit()
        return True

    def odf(self, section, ds, st):
        """The ODF connector next to each RTU: the first event after the launch on that end's own
        trace (the same rule as fms_pull.build_odf and the report's ODF Connection sheet).
        Pass at or under the ODF limit (default 0.5 dB at 1550), fail above it."""
        con = self.db(section)
        info = dict(con.execute("SELECT * FROM datasets WHERE id=?", (ds,)).fetchone())
        lim = float(st.get("odf_limit", 0.5))
        fib = {r["fno"]: dict(r) for r in con.execute("SELECT fno, ribbon, date_a, date_b FROM fibres WHERE ds=?", (ds,))}
        rows = {}
        for r in con.execute("SELECT * FROM odf WHERE ds=?", (ds,)):
            d = dict(r)
            v = d["w1550"]
            d["fail"] = v is not None and v > lim
            rows.setdefault(d["fno"], {})[d["end_"]] = d
        out, counts = [], {"A": {"tested": 0, "fail": 0, "none": 0}, "B": {"tested": 0, "fail": 0, "none": 0}}
        for fno in sorted(set(fib) | set(rows)):
            rec = {"fno": fno, "ribbon": ribbon_of(fno)}
            for end in ("A", "B"):
                e = (rows.get(fno) or {}).get(end)
                tested = bool((fib.get(fno) or {}).get("date_a" if end == "A" else "date_b"))
                if e and e["w1550"] is not None:
                    counts[end]["tested"] += 1
                    counts[end]["fail"] += e["fail"]
                elif tested:
                    counts[end]["none"] += 1
                rec[end] = {k: e[k] for k in ("pos", "type", "refl", "w1310", "w1550", "w1625", "sw", "fail")} if e else None
            out.append(rec)
        vals = {end: [r[end]["w1550"] for r in out if r[end] and r[end]["w1550"] is not None] for end in ("A", "B")}
        for end in ("A", "B"):
            xs = sorted(vals[end])
            counts[end]["median"] = round(xs[len(xs) // 2], 3) if xs else None
            counts[end]["worst"] = round(xs[-1], 3) if xs else None
        return {"dataset": info, "limit": lim, "site_a": info.get("site_a"), "site_b": info.get("site_b"),
                "counts": counts, "fibres": out, "settings": st}

    # ---------------------------------------------------------------- snapshots and compare
    def delete_dataset(self, section, ds):
        """Take one pull off this cable. The store is copied first (store.before-remove-<id>-<time>.sqlite);
        the raw results stay, so the same tests can be built again by a pull. Names, snapshots, assets and
        remediations are kept."""
        con = self.db(section)
        row = con.execute("SELECT id, label FROM datasets WHERE id=?", (ds,)).fetchone()
        if not row:
            raise KeyError("dataset")
        path = os.path.join(self.folder(section), "store.sqlite")
        con.commit()
        bak = os.path.join(self.folder(section), f"store.before-remove-{ds}-{time.strftime('%Y%m%d-%H%M%S')}.sqlite")
        dst = sqlite3.connect(bak)
        con.backup(dst)
        dst.close()
        for t in ("events", "fibres", "locations", "odf"):
            con.execute(f"DELETE FROM {t} WHERE ds=?", (ds,))
        con.execute("DELETE FROM datasets WHERE id=?", (ds,))
        con.commit()
        return {"removed": {"id": row["id"], "label": row["label"]}, "backup": os.path.basename(bak),
                "datasets": self.datasets(section)}

    def set_dataset(self, section, ds, label=None, baseline=None):
        """Rename a pull or mark it as the baseline. Results themselves are never changed."""
        con = self.db(section)
        row = con.execute("SELECT notes FROM datasets WHERE id=?", (ds,)).fetchone()
        if not row:
            raise KeyError("dataset")
        if label is not None and str(label).strip():
            con.execute("UPDATE datasets SET label=? WHERE id=?", (str(label).strip()[:80], ds))
        if baseline is not None:
            for r in con.execute("SELECT id, notes FROM datasets").fetchall():
                try:
                    n = json.loads(r["notes"] or "{}")
                except ValueError:
                    n = {}
                want = bool(baseline) and r["id"] == ds
                if bool(n.get("baseline")) != want:
                    n["baseline"] = want
                    con.execute("UPDATE datasets SET notes=? WHERE id=?", (json.dumps(n), r["id"]))
        con.commit()
        return self.datasets(section)

    def judged(self, section, ds, st):
        """Every splice judged in one pull, worst per fibre per location."""
        con = self.db(section)
        an = self.analysis(section, ds, st)
        meta = {l["loc"]: l for l in an["locs"]}
        groups, ravg = self._groups(con, ds)
        out = {}
        for key, grp in groups.items():
            m = meta.get(key[1])
            lim = self._limit(m, key[0], st, an["nr"])
            for e in grp:
                v, how, gain = splice_value(e["ab"], e["ba"], st, *ravg[key])
                if v is None or how == "one_way":
                    continue
                k = (e["fno"], e["loc"])
                if k not in out or v > out[k]["value"]:
                    out[k] = {"fno": e["fno"], "ribbon": e["ribbon"], "loc": e["loc"], "pos": e["pos"],
                              "value": round(v, 3), "how": how, "ab": e["ab"], "ba": e["ba"],
                              "limit": lim, "fail": v >= lim}
        return out, an

    def compare(self, section, base, new, st, group=None):
        """What changed between two pulls: solved, remaining, new and not seen now."""
        con = self.db(section)
        info = {d["id"]: d for d in self.datasets(section)}
        if base not in info or new not in info:
            raise KeyError("dataset")
        jb, anb = self.judged(section, base, st)
        jn, ann = self.judged(section, new, st)
        gsel = next((g for g in ann["groups"] if g["id"] == group), None)
        keep = gsel["fnos"] if gsel else None
        meta = {l["loc"]: l for l in anb["locs"]}
        meta.update({l["loc"]: l for l in ann["locs"]})
        by_fno = {}
        for r in jn.values():
            by_fno.setdefault(r["fno"], []).append(r)
        used, rows = set(), []

        def match(r):
            k = (r["fno"], r["loc"])
            if k in jn:
                return jn[k]
            near = [x for x in by_fno.get(r["fno"], []) if x["pos"] is not None and r["pos"] is not None
                    and abs(x["pos"] - r["pos"]) <= MATCH_TOL_M and (x["fno"], x["loc"]) not in used]
            return min(near, key=lambda x: abs(x["pos"] - r["pos"])) if near else None

        for r in jb.values():
            if keep is not None and r["fno"] not in keep:
                continue
            m = match(r)
            if m:
                used.add((m["fno"], m["loc"]))
            if not r["fail"] and not (m and m["fail"]):
                continue
            if r["fail"] and m is None:
                status = "not_seen"
            elif r["fail"] and m["fail"]:
                status = "remaining"
            elif r["fail"]:
                status = "solved"
            else:
                status = "new"
            rows.append(self._crow(r, m, status, meta))
        for k, m in jn.items():
            if k in used or not m["fail"] or (keep is not None and m["fno"] not in keep):
                continue
            rows.append(self._crow(None, m, "new", meta))
        statuses = ("solved", "remaining", "new", "not_seen")
        totals = {s: sum(1 for r in rows if r["status"] == s) for s in statuses}
        totals["fail_before"] = sum(1 for r in jb.values() if r["fail"] and (keep is None or r["fno"] in keep))
        totals["fail_now"] = sum(1 for r in jn.values() if r["fail"] and (keep is None or r["fno"] in keep))
        locs = {}
        for r in rows:
            L = locs.setdefault(r["loc"], {"loc": r["loc"], "pos": r["loc_pos"], "spur": r["spur"],
                                           "sodf": r["sodf"], **{s: 0 for s in statuses}})
            L[r["status"]] += 1
        ribs = {}
        for r in rows:
            R = ribs.setdefault(r["ribbon"], {"ribbon": r["ribbon"], **{s: 0 for s in statuses}})
            R[r["status"]] += 1
        rows.sort(key=lambda r: (statuses.index(r["status"]), r["loc_pos"] if r["loc_pos"] is not None else 9e9, r["fno"]))
        return {"base": info[base], "new": info[new], "settings": st, "totals": totals,
                "locations": sorted(locs.values(), key=lambda x: x["pos"] if x["pos"] is not None else 9e9),
                "ribbons": sorted(ribs.values(), key=lambda x: x["ribbon"]),
                "rows": rows,
                "groups": [{k: v for k, v in g.items() if k != "fnos"} for g in ann["groups"]],
                "group": group if gsel else None,
                "spurs": [sp["name"] for sp in ann["spurs"]]}

    # ---------------------------------------------------------------- history: pulls kept over time
    TABLES = ("fibres", "events", "locations", "odf")

    def export_pull(self, section, ds):
        """One pull as a small gzip JSON, to keep it beyond the work folder (the relay's log repo).
        The results themselves (both ends, every event) go with it, so it can be judged again with
        any settings and compared with any other pull."""
        con = self.db(section)
        d = con.execute("SELECT * FROM datasets WHERE id=?", (ds,)).fetchone()
        if not d:
            raise KeyError("dataset")
        out = {"v": 1, "dataset": {k: d[k] for k in d.keys() if k != "id"}}
        for t in self.TABLES:
            cols = [r[1] for r in con.execute(f"PRAGMA table_info({t})") if r[1] != "ds"]
            out[t] = {"cols": cols, "rows": [list(r) for r in con.execute(
                f"SELECT {','.join(cols)} FROM {t} WHERE ds=?", (ds,))]}
        return gzip.compress(json.dumps(out, separators=(",", ":")).encode("utf-8"))

    def import_pull(self, section, blob):
        """A pull kept earlier, back into the store. Returns its new id, or None when the same
        results are already here."""
        data = json.loads(gzip.decompress(blob).decode("utf-8"))
        d = data["dataset"]
        con = self.db(section)
        if d.get("content_sig") and con.execute("SELECT 1 FROM datasets WHERE content_sig=?", (d["content_sig"],)).fetchone():
            return None
        have = {r[1] for r in con.execute("PRAGMA table_info(datasets)")} - {"id"}
        keys = [k for k in d if k in have]
        cur = con.execute(f"INSERT INTO datasets({','.join(keys)}) VALUES({','.join('?' * len(keys))})",
                          [d[k] for k in keys])
        ds = cur.lastrowid
        for t in self.TABLES:
            part = data.get(t) or {}
            cols = part.get("cols") or []
            tcols = {r[1] for r in con.execute(f"PRAGMA table_info({t})")}
            use = [i for i, c in enumerate(cols) if c in tcols]
            if not use:
                continue
            names = ["ds"] + [cols[i] for i in use]
            con.executemany(f"INSERT INTO {t}({','.join(names)}) VALUES({','.join('?' * len(names))})",
                            [[ds] + [r[i] for i in use] for r in part.get("rows") or []])
        con.commit()
        con.close()
        self.db(section).close()                 # realigns it when it was kept by an older version
        return ds

    # ---------------------------------------------------------------- snapshots: checked tests
    def snapshots(self, section):
        """Tests that have been checked and saved as snapshots, each with the pass/fail criteria it was
        saved with (EXFO style), keyed by the test's content signature so they survive a rebuild."""
        try:
            with open(os.path.join(self.folder(section), "snapshots.json"), encoding="utf-8") as fh:
                return [x for x in json.load(fh).get("snapshots", []) if x.get("sig")]
        except (OSError, ValueError):
            return []

    def save_snapshots(self, section, snaps):
        os.makedirs(self.folder(section), exist_ok=True)
        path = os.path.join(self.folder(section), "snapshots.json")
        with open(path + ".tmp", "w", encoding="utf-8") as fh:
            json.dump({"snapshots": snaps}, fh, indent=1)
        os.replace(path + ".tmp", path)
        if SNAPS_SAVED:
            try:
                SNAPS_SAVED(section, snaps)
            except Exception:               # noqa: BLE001
                pass

    def snapshot_of(self, section, ds):
        d = next((x for x in self.datasets(section) if x["id"] == ds), None)
        if not d:
            return None
        return next((x for x in self.snapshots(section) if x["sig"] == d.get("content_sig")), None)

    def lineup(self, section, base, new, st):
        """Does one test line up with another before they are compared: the same joints at the same
        places. Joints are matched by name, or by place when unnamed; small clusters are left out."""
        a, b = self.analysis(section, base, st), self.analysis(section, new, st)
        def big(an):
            n = max(3, round(0.05 * len(an["fibres"])))
            return [l for l in an["locs"] if _is_named(l) or len(l["fnos"]) >= n]
        la, lb = big(a), big(b)
        joints = []
        for l in la:
            joints.append({"loc": l["loc"], "pos": l["pos"], "spur_d": l.get("spur_d"), "named": _is_named(l),
                           "a": l, "b": None})
        matched, moved, only_new = 0, [], []
        used = set()
        for l in lb:
            j = _trend_joint([x for x in joints if x["b"] is None], l, used)
            if j is None:
                only_new.append(_lj(l)); continue
            j["b"] = l; used.add(id(j)); matched += 1
            shift = l["pos"] - j["pos"]
            if abs(shift) > 10 or (j["loc"] != l["loc"] and _is_named(l) and j["named"]):
                moved.append({"before": j["loc"], "now": l["loc"], "pos_before": round(j["pos"], 1),
                              "pos_now": round(l["pos"], 1), "shift_m": round(shift, 1)})
        only_prev = [_lj(j["a"]) for j in joints if j["b"] is None]
        return {"base": base, "new": new, "matched": matched, "moved": moved, "only_before": only_prev,
                "only_now": only_new, "ok": not moved and not only_prev and not only_new}

    def review(self, section, ds, st):
        """Everything to check before a test is saved as a snapshot: the ODF at each end, the ODFs along
        the route, the joints and spurs found, the pass/fail criteria, and how it lines up with the
        snapshot before it."""
        info = next((x for x in self.datasets(section) if x["id"] == ds), None)
        if not info:
            raise KeyError("dataset")
        an = self.analysis(section, ds, st)
        od = self.odf(section, ds, st)
        ends = {}
        for end, site in (("A", od.get("site_a")), ("B", od.get("site_b"))):
            pos = sorted(f[end]["pos"] for f in od["fibres"] if f.get(end) and f[end].get("pos") is not None)
            ends[end] = {"site": site or ("End " + end), "read": od["counts"][end]["tested"],
                         "none": od["counts"][end]["none"], "fail": od["counts"][end]["fail"],
                         "median": od["counts"][end]["median"], "worst": od["counts"][end]["worst"],
                         "pos_m": round(pos[len(pos) // 2], 1) if pos else None}
        nmin = max(3, round(0.05 * len(an["fibres"])))
        route = [{"loc": l["loc"], "pos": l["pos"], "kind": l.get("kind") or "", "spur": l.get("spur"),
                  "spur_leg": l.get("spur_leg"), "spur_m": l.get("spur_m"), "n_fibres": len(l["fnos"]),
                  "named": _is_named(l), "scheduled": l.get("scheduled")}
                 for l in an["locs"] if _is_odf_loc(l)]
        far = []                            # unnamed spur ends that are probably the REB ODF
        for sp in an["spurs"]:
            ends_ = [l for l in an["locs"] if l["loc"] in sp["locations"] and l.get("spur_leg") == "end"]
            for l in ends_:
                if not _is_odf_loc(l):
                    far.append({"loc": l["loc"], "pos": l["pos"], "spur": sp["name"], "spur_m": l.get("spur_m"),
                                "n_fibres": len(l["fnos"]), "named": _is_named(l), "scheduled": l.get("scheduled")})
        big = [l for l in an["locs"] if len(l["fnos"]) >= nmin or _is_named(l)]
        snaps = {x["sig"]: x for x in self.snapshots(section)}
        allds = sorted(self.datasets(section), key=lambda d: (d["t_to"] or "", d["id"]))
        prev = [d for d in allds if d.get("content_sig") in snaps and d["id"] != ds and (d["t_to"] or "") <= (info["t_to"] or "")]
        line = None
        if prev:
            line = self.lineup(section, prev[-1]["id"], ds, st)
            line["base_label"] = prev[-1]["label"]
        return {"dataset": {k: info.get(k) for k in ("id", "label", "t_from", "t_to", "n_fibres", "both_ends", "one_end",
                                                     "schedule", "content_sig")},
                "snapshot": snaps.get(info.get("content_sig")), "settings": st, "ends": ends,
                "route_odfs": route, "spur_ends": far,
                "spurs": [{k: sp.get(k) for k in ("name", "start", "length", "ribbons", "locations")} for sp in an["spurs"]],
                "joints": {"all": len(big), "named": sum(1 for l in big if _is_named(l)),
                           "unnamed": [{"loc": l["loc"], "pos": l["pos"], "n_fibres": len(l["fnos"])}
                                       for l in big if not _is_named(l)][:40]},
                "groups": [{k: g.get(k) for k in ("id", "label")} for g in an["groups"]],
                "lineup": line}

    def set_snapshot(self, section, ds, st, by="", note="", remove=False, checks=None):
        """Save a checked test as a snapshot with the criteria it was checked with, or take it off."""
        info = next((x for x in self.datasets(section) if x["id"] == ds), None)
        if not info:
            raise KeyError("dataset")
        snaps = [x for x in self.snapshots(section) if x["sig"] != info.get("content_sig")]
        entry = None
        if not remove:
            rv = self.review(section, ds, st)
            entry = {"sig": info.get("content_sig"), "label": info["label"], "t_from": info["t_from"],
                     "t_to": info["t_to"], "by": by, "at": time.strftime("%Y-%m-%d %H:%M"),
                     "note": " ".join(str(note or "").split())[:300], "settings": dict(st),
                     "checks": checks or {}, "ends": rv["ends"],
                     "route_odfs": [r["loc"] for r in rv["route_odfs"]], "spurs": [sp["name"] for sp in rv["spurs"]],
                     "joints": rv["joints"]["all"], "named": rv["joints"]["named"],
                     "lineup": None if not rv["lineup"] else {
                         "with": rv["lineup"]["base_label"], "matched": rv["lineup"]["matched"],
                         "moved": len(rv["lineup"]["moved"]), "only_before": len(rv["lineup"]["only_before"]),
                         "only_now": len(rv["lineup"]["only_now"])}}
            snaps.append(entry)
        snaps.sort(key=lambda x: (x.get("t_to") or "", x.get("at") or ""))
        self.save_snapshots(section, snaps)
        return {"snapshot": entry, "snapshots": snaps}

    def remedials(self, section, ds, st, group=None):
        """The remedial list for one test: splice remedials (fails at joints, worst splice per fibre per
        joint, with the S-ODF limit where it applies) and, kept apart, ODF remedials (the ODF at each end
        over the ODF limit, on that end's own trace)."""
        info = next((x for x in self.datasets(section) if x["id"] == ds), None)
        if not info:
            raise KeyError("dataset")
        rows, an = self.judged(section, ds, st)
        gsel = next((g for g in an["groups"] if g["id"] == group), None)
        keep = gsel["fnos"] if gsel else None
        meta = {l["loc"]: l for l in an["locs"]}
        sp = []
        for r in rows.values():
            if not r["fail"] or (keep is not None and r["fno"] not in keep):
                continue
            m = meta.get(r["loc"]) or {}
            sp.append({"fno": r["fno"], "ribbon": r["ribbon"], "loc": r["loc"], "pos": m.get("pos"), "kind": m.get("kind") or "",
                       "spur": m.get("spur"), "spur_leg": m.get("spur_leg"), "spur_m": m.get("spur_m"),
                       "sodf": bool(m.get("sodf")), "ab": r["ab"], "ba": r["ba"], "value": r["value"], "how": r["how"],
                       "limit": r["limit"], "trace_pos": r["pos"]})
        sp.sort(key=lambda r: (r["pos"] if r["pos"] is not None else 9e9, r["fno"]))
        o = self.odf(section, ds, st)
        od = []
        for f in o["fibres"]:
            if keep is not None and f["fno"] not in keep:
                continue
            for end in ("A", "B"):
                x = f.get(end)
                if x and x.get("fail"):
                    od.append({"fno": f["fno"], "ribbon": f["ribbon"], "end": end,
                               "site": o.get("site_a" if end == "A" else "site_b") or ("End " + end),
                               "pos": x.get("pos"), "type": x.get("type"), "refl": x.get("refl"), "w1310": x.get("w1310"),
                               "w1550": x.get("w1550"), "w1625": x.get("w1625"), "limit": o["limit"]})
        od.sort(key=lambda r: (r["end"], r["fno"]))
        return {"dataset": {k: info.get(k) for k in ("id", "label", "t_from", "t_to", "n_fibres", "site_a", "site_b")},
                "settings": st, "group": group if gsel else None, "splices": sp, "odfs": od,
                "counts": {"splice": len(sp), "odf": len(od), "odf_A": sum(1 for r in od if r["end"] == "A"),
                           "odf_B": sum(1 for r in od if r["end"] == "B"),
                           "joints": len({r["loc"] for r in sp}), "fibres": len({r["fno"] for r in sp} | {r["fno"] for r in od})}}

    def trend(self, section, st, group=None, ids=None):
        """Fails at every joint in every kept pull (or the chosen ones), oldest first."""
        t = self.track(section, ids, st, group)
        return {k: t[k] for k in ("pulls", "joints", "settings", "group")}

    def track(self, section, ids, st, group=None, frozen=True):
        """Brunel 1.1.3: follow every splice across as many pulls as wanted (all when ids is empty),
        oldest first: its loss in each test and what happened to it (still failing, new, fixed, came
        back, not in the latest). Joints are matched across pulls by name, or by place when unnamed;
        a splice is one fibre at one joint (its worst splice there, as in compare)."""
        allp = sorted(self.datasets(section), key=lambda d: (d["t_to"] or "", d["id"]))
        want = {int(i) for i in (ids or [])}
        pulls = [d for d in allp if not want or d["id"] in want]
        joints, out_pulls, sp = [], [], {}
        snaps = {x["sig"]: x for x in self.snapshots(section)} if frozen else {}
        for d in pulls:
            sn = snaps.get(d.get("content_sig"))
            dst = dict(DEFAULT_SETTINGS) | sn["settings"] if sn and sn.get("settings") else st
            rows, an = self.judged(section, d["id"], dst)
            gsel = next((g for g in an["groups"] if g["id"] == group), None)
            keep = gsel["fnos"] if gsel else None
            fails, judged = {}, {}
            for r in rows.values():
                if keep is not None and r["fno"] not in keep:
                    continue
                judged[r["loc"]] = judged.get(r["loc"], 0) + 1
                if r["fail"]:
                    fails[r["loc"]] = fails.get(r["loc"], 0) + 1
            out_pulls.append({"id": d["id"], "label": d["label"], "t_from": d["t_from"], "t_to": d["t_to"],
                              "fails": sum(fails.values()), "judged": sum(judged.values()),
                              "baseline": bool((d.get("notes") or {}).get("baseline")),
                              "snapshot": bool(sn), "settings": dst if sn else None})
            used, jof = set(), {}
            for l in an["locs"]:
                if l["loc"] not in judged:
                    continue
                j = _trend_joint(joints, l, used)
                if j is None:
                    j = {"key": len(joints), "loc": l["loc"], "pos": l["pos"], "spur_d": l.get("spur_d"),
                         "spur_leg": l.get("spur_leg"), "spur_m": l.get("spur_m"), "spur": l.get("spur"),
                         "sodf": l.get("sodf"), "kind": l.get("kind"), "named": _is_named(l), "fails": {}, "judged": {}}
                    joints.append(j)
                elif _is_named(l) and not j["named"]:
                    j.update(loc=l["loc"], named=True, kind=l.get("kind"))
                used.add(id(j))
                jof[l["loc"]] = j
                j["fails"][d["id"]] = j["fails"].get(d["id"], 0) + fails.get(l["loc"], 0)
                j["judged"][d["id"]] = j["judged"].get(d["id"], 0) + judged[l["loc"]]
                j["spur"] = l.get("spur") or j["spur"]
            for r in rows.values():
                if keep is not None and r["fno"] not in keep:
                    continue
                j = jof.get(r["loc"])
                if j is None:
                    continue
                k = (r["fno"], j["key"])
                e = sp.setdefault(k, {"fno": r["fno"], "ribbon": r["ribbon"], "joint": j, "values": {}, "fail": {},
                                      "how": {}, "limit": {}})
                old = e["values"].get(d["id"])
                if old is None or r["value"] > old:
                    e["values"][d["id"]], e["fail"][d["id"]] = r["value"], r["fail"]
                    e["how"][d["id"]], e["limit"][d["id"]] = r["how"], r["limit"]
        order = [p["id"] for p in out_pulls]
        lab = {p["id"]: p["label"] for p in out_pulls}
        rows, counts = [], {k: 0 for k in TRACK_STATUS}
        for e in sp.values():
            if not any(e["fail"].values()):
                continue
            status, since = _track_status([e["fail"].get(i) for i in order], order)
            counts[status] += 1
            j = e["joint"]
            vals = [e["values"].get(i) for i in order]
            seen = [v for v in vals if v is not None]
            rows.append({"fno": e["fno"], "ribbon": e["ribbon"], "loc": j["loc"], "pos": j["pos"],
                         "spur": j["spur"], "spur_leg": j["spur_leg"], "spur_m": j["spur_m"], "sodf": j["sodf"],
                         "values": {str(i): e["values"].get(i) for i in order},
                         "fail": {str(i): e["fail"].get(i) for i in order},
                         "how": {str(i): e["how"].get(i) for i in order},
                         "limit": next((e["limit"][i] for i in reversed(order) if i in e["limit"]), None),
                         "status": status, "since": since, "since_label": lab.get(since),
                         "change": round(seen[-1] - seen[0], 3) if len(seen) >= 2 else None,
                         "fails_in": sum(1 for v in e["fail"].values() if v)})
        rows.sort(key=lambda r: (TRACK_STATUS.index(r["status"]), r["pos"], r["fno"]))
        joints.sort(key=lambda j: (j["pos"], -1 if j["spur_d"] is None else j["spur_d"]))
        for j in joints:
            j.pop("named", None); j.pop("key", None)
        # ODF remedials, kept apart from splices: the ODF at each end on that end's own trace, at the ODF limit
        od_rows, od_counts = [], {k: 0 for k in TRACK_STATUS}
        ods, sites = {}, {}
        for d in pulls:
            sn = snaps.get(d.get("content_sig"))
            dst = dict(DEFAULT_SETTINGS) | sn["settings"] if sn and sn.get("settings") else st
            o = self.odf(section, d["id"], dst)
            sites = {"A": o.get("site_a") or "End A", "B": o.get("site_b") or "End B"}
            for f in o["fibres"]:
                for end in ("A", "B"):
                    x = f.get(end)
                    if not x or x.get("w1550") is None:
                        continue
                    e = ods.setdefault((f["fno"], end), {"values": {}, "fail": {}, "pos": {}})
                    e["values"][d["id"]], e["fail"][d["id"]], e["pos"][d["id"]] = x["w1550"], bool(x["fail"]), x.get("pos")
        for (fno, end), e in ods.items():
            if not any(e["fail"].values()):
                continue
            status, since = _track_status([e["fail"].get(i) for i in order], order)
            od_counts[status] += 1
            seen = [e["values"][i] for i in order if i in e["values"]]
            od_rows.append({"fno": fno, "ribbon": ribbon_of(fno), "end": end, "site": sites.get(end),
                            "pos": next((e["pos"][i] for i in reversed(order) if i in e["pos"]), None),
                            "values": {str(i): e["values"].get(i) for i in order},
                            "fail": {str(i): e["fail"].get(i) for i in order}, "status": status, "since": since,
                            "since_label": lab.get(since), "change": round(seen[-1] - seen[0], 3) if len(seen) >= 2 else None,
                            "fails_in": sum(1 for v in e["fail"].values() if v)})
        od_rows.sort(key=lambda r: (TRACK_STATUS.index(r["status"]), r["end"], r["fno"]))
        for p_, d in zip(out_pulls, pulls):
            p_["odf_fails"] = sum(1 for e in ods.values() if e["fail"].get(d["id"]))
        return {"pulls": out_pulls, "joints": joints, "rows": rows, "counts": counts, "settings": st,
                "odf_rows": od_rows, "odf_counts": od_counts,
                "group": group, "all": [{"id": d["id"], "label": d["label"], "t_to": d["t_to"]} for d in allp]}

    @staticmethod
    def _crow(b, n, status, meta):
        r = n or b
        loc = r["loc"]
        m = meta.get(loc) or {}
        return {"fno": r["fno"], "ribbon": r["ribbon"], "loc": loc,
                "loc_pos": m.get("pos", r["pos"]), "spur": m.get("spur"), "sodf": bool(m.get("sodf")),
                "status": status,
                "before": b["value"] if b else None, "before_how": b["how"] if b else None,
                "now": n["value"] if n else None, "now_how": n["how"] if n else None,
                "limit": r["limit"],
                "change": round(n["value"] - b["value"], 3) if (b and n) else None}


# --------------------------------------------------------------------------- analysis helpers
def parse_ribbons(spec):
    out = set()
    for part in str(spec or "").replace(" ", "").split(","):
        if not part:
            continue
        try:
            if "-" in part:
                a, b = part.split("-", 1)
                out.update(range(int(a), int(b) + 1))
            else:
                out.add(int(part))
        except ValueError:
            pass
    return out


def suggest_nr(groups, meta, min_gap=0.02):
    """Which ribbons carry two splices at the S-ODFs: average both end loss per ribbon at
    every S-ODF, then split the ribbons at the biggest step between those averages."""
    per = {}
    for (rb, loc), grp in groups.items():
        if not (meta.get(loc) or {}).get("sodf"):
            continue
        for e in grp:
            if e["ab"] is not None and e["ba"] is not None:
                per.setdefault(rb, []).append((e["ab"] + e["ba"]) / 2)
    means = {rb: sum(v) / len(v) for rb, v in per.items() if len(v) >= 3}
    if len(means) < 4:
        return None
    order = sorted(means)                       # NR ribbons sit in one run, so split by ribbon order
    vals = [means[r] for r in order]
    tot = sum(vals)
    best = None
    for k in range(1, len(order)):
        lo, hi = vals[:k], vals[k:]
        ml, mh = sum(lo) / len(lo), (tot - sum(lo)) / len(hi)
        score = len(lo) * len(hi) * (ml - mh) ** 2           # between class variance
        if best is None or score > best[0]:
            best = (score, k, ml, mh)
    _, k, ml, mh = best
    first = ml > mh
    high = order[:k] if first else order[k:]
    diff = abs(ml - mh)
    out = {"diff": round(diff, 3), "high_avg": round(max(ml, mh), 3), "low_avg": round(min(ml, mh), 3),
           "means": {r: round(m, 3) for r, m in sorted(means.items())}}
    if diff < min_gap:
        return out | {"spec": "", "ribbons": []}
    if first and high[0] == order[0]:
        high = list(range(1, high[-1] + 1))   # ribbons before the first S-ODF reading belong with the run
    return out | {"spec": ribbon_span(high).replace("R", "").replace(" ", ""), "ribbons": high}


def is_sodf(kind):
    return bool(kind) and re.sub(r"[^A-Z]", "", str(kind).upper()) == "SODF"


def ribbon_span(ribbons):
    rs = sorted(ribbons)
    if not rs:
        return ""
    parts, start, prev = [], rs[0], rs[0]
    for r in rs[1:] + [None]:
        if r is not None and r == prev + 1:
            prev = r
            continue
        parts.append(f"R{start}" if start == prev else f"R{start}-{prev}")
        if r is not None:
            start = prev = r
    return ", ".join(parts)


def fibre_groups(fibres, tol=100.0):
    """Group fibres by measured length: fibres whose lengths chain within tol metres
    take the same route. A fibre with no length joins the group of its ribbon."""
    pts = sorted(((f["len_a"] or f["len_b"] or 0), f["fno"]) for f in fibres if (f["len_a"] or f["len_b"]))
    runs, cur = [], []
    for p in pts:
        if cur and p[0] - cur[-1][0] > tol:
            runs.append(cur); cur = []
        cur.append(p)
    if cur:
        runs.append(cur)
    big = [r for r in runs if len(r) >= 3]
    small = [p for r in runs if len(r) < 3 for p in r]
    groups = [{"fnos": {f for _, f in r}, "km": statistics.median(x for x, _ in r) / 1000.0} for r in big]
    by_ribbon = {}
    for g in groups:
        for f in g["fnos"]:
            by_ribbon.setdefault(ribbon_of(f), []).append(g)
    lost = [f["fno"] for f in fibres if not (f["len_a"] or f["len_b"])] + [f for _, f in small]
    for f in lost:
        cand = by_ribbon.get(ribbon_of(f)) or groups
        if cand:
            max(cand, key=lambda g: len(g["fnos"]))["fnos"].add(f)
    groups.sort(key=lambda g: g["km"])
    for i, g in enumerate(groups, 1):
        rb = {ribbon_of(f) for f in g["fnos"]}
        g.update(id=f"G{i}", n=len(g["fnos"]), ribbons=ribbon_span(rb), ribbon_list=sorted(rb),
                 label=f"Group {i} · {g['km']:.1f} km · {ribbon_span(rb)} ({len(g['fnos'])} fibres)")
    return groups


def refine_groups(groups, loc_f, hi=0.8, lo=0.2, min_joints=2):
    """Second pass on the length groups: split a group when some of its ribbons clearly go
    through joints the others clearly do not (seen by at least 80% of one ribbon's fibres and
    at most 20% of the other's, on at least 2 joints). Single missing splices are normal, as low
    loss splices are not always detected, so they never split a group."""
    out = []
    for g in groups:
        by_rib = {}
        for f in g["fnos"]:
            by_rib.setdefault(ribbon_of(f), set()).add(f)
        ribs = sorted(by_rib)
        frac = {r: {loc: len(fs & by_rib[r]) / len(by_rib[r]) for loc, fs in loc_f.items()} for r in ribs}

        def differ(a, b):
            return [l for l in loc_f if (frac[a][l] >= hi and frac[b][l] <= lo) or (frac[b][l] >= hi and frac[a][l] <= lo)]
        parent = {r: r for r in ribs}

        def find(r):
            while parent[r] != r:
                parent[r] = parent[parent[r]]
                r = parent[r]
            return r
        for i, a in enumerate(ribs):
            for b in ribs[i + 1:]:
                if len(differ(a, b)) < min_joints:
                    parent[find(a)] = find(b)
        comps = {}
        for r in ribs:
            comps.setdefault(find(r), []).append(r)
        parts = sorted(comps.values(), key=lambda x: x[0])
        if len(parts) == 1:
            out.append(g | {"route": "confirmed by splice locations"})
            continue
        for k, rs in enumerate(parts):
            fn = set().union(*(by_rib[r] for r in rs))
            d = differ(rs[0], next(x for p in parts if p is not rs for x in p))
            sub = g | {"id": f"{g['id']}{'abcdefghij'[k]}", "fnos": fn, "n": len(fn), "ribbon_list": rs,
                       "ribbons": ribbon_span(rs), "route": f"split by joints: {', '.join(d[:3])}"}
            sub["label"] = f"Group {sub['id'][1:]} · {g['km']:.1f} km · {sub['ribbons']} ({len(fn)} fibres)"
            out.append(sub)
    return out


def aligned_spurs(locs, loc_f, groups):
    """Spurs straight from the backbone alignment: the spur locations that leave the backbone
    at the same junction make one spur, named after its REB or ODF."""
    gid = {f: g["id"] for g in groups for f in g["fnos"]}
    sp = [l for l in locs if l.get("spur_d") is not None]
    out, cur = [], []
    for l in sorted(sp, key=lambda x: (x["pos"], x["spur_d"])):
        if cur and l["pos"] - cur[-1]["pos"] > 60:
            out.append(cur); cur = []
        cur.append(l)
    if cur:
        out.append(cur)
    spurs = []
    for run in out:
        run.sort(key=lambda x: x["spur_d"])
        far = max(run, key=lambda x: (ci_align._turn_name(x["loc"]) or ci_align._turn_name(x.get("kind")),
                                      -abs(x["spur_d"] - (x["spur_len"] or 0))))
        fn = set().union(*(loc_f.get(l["loc"], set()) for l in run))
        cnt = {}
        for f in fn:
            if f in gid:
                cnt[gid[f]] = cnt.get(gid[f], 0) + 1
        jpos = statistics.median(l["pos"] for l in run)
        nm = (f"Spur at {jpos / 1000:.2f} km" if "(unscheduled)" in far["loc"] or far["loc"].startswith("Unnamed ")
              else re.sub(r"\s+(N|S)?ODF$", "", far["loc"]) + " spur")
        spurs.append({"name": nm,
                      "group": max(cnt, key=cnt.get) if cnt else None,
                      "locations": [l["loc"] for l in run], "start": jpos, "end": jpos, "junction": jpos,
                      "length": round(max(l["spur_len"] or 0 for l in run)),
                      "ribbons": ribbon_span({ribbon_of(f) for f in fn})})
    return spurs


def detect_spurs(locs, loc_f, groups, total, gap=1500.0):
    """A spur shows as a run of locations, a short distance apart, seen almost only by a
    small group of fibres (the ribbons that leave the main line to a REB and come back).
    Small clusters inside the run are allowed; a location most fibres see ends it."""
    gid = {f: g["id"] for g in groups for f in g["fnos"]}
    gsize = {g["id"]: len(g["fnos"]) for g in groups}
    marks = []
    for l in locs:
        f = [x for x in loc_f.get(l["loc"], ()) if x in gid]
        if len(f) < 3:
            marks.append(("minor", None)); continue
        cnt = {}
        for x in f:
            cnt[gid[x]] = cnt.get(gid[x], 0) + 1
        g, n = max(cnt.items(), key=lambda kv: kv[1])
        if n / len(f) >= 0.8 and gsize[g] <= total / 3 and len(f) >= 0.25 * gsize[g]:
            marks.append(("spur", g))
        elif len(f) < 0.10 * total:
            marks.append(("minor", None))
        else:
            marks.append(("main", None))
    spurs, run = [], []

    def close():
        while run and run[-1][1] != "spur":
            run.pop()
        if run:
            members = [l for l, k, _ in run if k == "spur"]
            odf = next((l for l in members if (l.get("kind") or "").upper().endswith("ODF")), None)
            name = (odf["loc"] if odf else members[0]["loc"]) + " spur"
            spurs.append({"name": name, "group": run[0][2], "locations": [l["loc"] for l, _, _ in run],
                          "start": run[0][0]["pos"], "end": run[-1][0]["pos"]})
    for l, (kind, g) in zip(locs, marks):
        if kind == "spur":
            if run and (g != run[0][2] or l["pos"] - run[-1][0]["pos"] > gap):
                close(); run = []
            run.append((l, "spur", g))
        elif kind == "minor" and run and l["pos"] - run[-1][0]["pos"] <= gap:
            run.append((l, "minor", run[0][2]))
        else:
            close(); run = []
    close()
    return spurs
